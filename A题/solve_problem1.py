# -*- coding: utf-8 -*-
"""
2026 高教社杯国赛 A题「药材的烘干问题」
问题1：预热平衡阶段 药材温度与水分浓度变化规律
模型：一维圆柱径向 耦合传热传质，显式有限差分
"""
import os
import numpy as np
import openpyxl

# ============ 附录2 物性参数 ============
rho = 820.0      # 密度 kg/m^3
cp  = 2600.0     # 比热容 J/(kg·K)
k   = 0.36       # 热传导系数 W/(m·K)
h   = 25.0       # 对流换热系数 W/(m^2·K)
hm  = 8e-7       # 对流传质系数 m/s
D0  = 7e-9       # 扩散系数前置系数 m^2/s


def D_of_C(C):
    """水分浓度扩散系数 D(C) = 7e-9 * exp(-0.89 C)"""
    return D0 * np.exp(-0.89 * C)


# ============ 几何与网格 ============
R    = 0.02              # 半径 2 cm = 0.02 m
N    = 20                # 20 等分 -> dr = 0.1 cm
dr   = R / N             # 0.001 m
r    = np.arange(N + 1) * dr          # 节点坐标 0 .. 0.02
rp   = r + dr / 2                     # 界面坐标 i+1/2
rm   = r - dr / 2                     # 界面坐标 i-1/2（rm[0] 用不到）

alpha = k / (rho * cp)                # 热扩散系数 m^2/s

# ============ 初始条件 ============
T0 = 28.0
C0 = 2.55

# ============ 读取附件1：烘房边界条件 ============
def load_oven_data():
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), '附件')
    f = os.path.join(base, '附件1.xlsx')
    wb = openpyxl.load_workbook(f, data_only=True)
    ws = wb['Sheet1']
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    t = np.array([d[0] for d in data], dtype=float)
    T = np.array([d[1] for d in data], dtype=float)
    C = np.array([d[2] for d in data], dtype=float)
    return t, T, C


t_oven, T_oven, C_oven = load_oven_data()


def T_inf(t):
    return np.interp(t, t_oven, T_oven)


def C_inf(t):
    return np.interp(t, t_oven, C_oven)


def run(dt=1.0, t_end=1800.0):
    T = np.full(N + 1, T0)
    C = np.full(N + 1, C0)
    n_steps = int(round(t_end / dt))
    T_hist = [T.copy()]
    C_hist = [C.copy()]
    t_now = 0.0
    for _ in range(n_steps):
        # 当前时刻边界条件
        global_t = t_now
        t_now += dt
        # 用本步结束时刻的边界值（半隐式可取中点，这里取显式端点）
        T_amb = T_inf(t_now)
        C_amb = C_inf(t_now)

        # 直接展开一步（含边界），传入边界值
        T, C = step_with_bc(T, C, dt, T_amb, C_amb)
        T_hist.append(T.copy())
        C_hist.append(C.copy())
    return np.array(T_hist), np.array(C_hist)


def step_with_bc(T, C, dt, T_amb, C_amb):
    Tn = T.copy()
    Cn = C.copy()
    Tn[0] = T[0] + dt * alpha * 4.0 * (T[1] - T[0]) / dr**2
    D0c = D_of_C(0.5 * (C[0] + C[1]))
    Cn[0] = C[0] + dt * 4.0 * D0c * (C[1] - C[0]) / dr**2

    i = np.arange(1, N)
    lapT = (rp[i] * (T[i + 1] - T[i]) - rm[i] * (T[i] - T[i - 1])) / (r[i] * dr**2)
    Tn[i] = T[i] + dt * alpha * lapT
    Dp = D_of_C(0.5 * (C[i] + C[i + 1]))
    Dm = D_of_C(0.5 * (C[i] + C[i - 1]))
    lapC = (rp[i] * Dp * (C[i + 1] - C[i]) - rm[i] * Dm * (C[i] - C[i - 1])) / (r[i] * dr**2)
    Cn[i] = C[i] + dt * lapC

    # 边界
    T_ghost = T[N - 1] - (2 * dr * h / k) * (T[N] - T_amb)
    lapT_N = (rp[N] * (T_ghost - T[N]) - rm[N] * (T[N] - T[N - 1])) / (r[N] * dr**2)
    Tn[N] = T[N] + dt * alpha * lapT_N

    DN = D_of_C(C[N])
    C_ghost = C[N - 1] - (2 * dr * hm / DN) * (C[N] - C_amb)
    DpN = D_of_C(0.5 * (C[N] + C_ghost))
    DmN = D_of_C(0.5 * (C[N] + C[N - 1]))
    lapC_N = (rp[N] * DpN * (C_ghost - C[N]) - rm[N] * DmN * (C[N] - C[N - 1])) / (r[N] * dr**2)
    Cn[N] = C[N] + dt * lapC_N

    return Tn, Cn


if __name__ == '__main__':
    dt = 1.0
    t_end = 1800.0
    T_hist, C_hist = run(dt, t_end)
    print('history shape:', T_hist.shape)   # (1801, 21)

    # ---- 表1 / 表2 要求：时间 100..1800 s，距离 0,0.5,1,1.5,2 cm ----
    times = [100, 300, 600, 900, 1200, 1500, 1800]
    dist_idx = [0, 5, 10, 15, 20]           # 对应 0,0.5,1,1.5,2 cm
    dist_cm = [0, 0.5, 1, 1.5, 2.0]

    print('\n表1  30分钟内药材的温度 (°C)')
    print('时间/s  ' + '  '.join(f'{d}cm' for d in dist_cm))
    for t in times:
        row = T_hist[t]                     # t 秒 == 索引 t（dt=1）
        print(f'{t:<7d}' + '  '.join(f'{row[j]:.4f}' for j in dist_idx))

    print('\n表2  30分钟内药材的水分浓度 (kg/kg)')
    print('时间/s  ' + '  '.join(f'{d}cm' for d in dist_cm))
    for t in times:
        row = C_hist[t]
        print(f'{t:<7d}' + '  '.join(f'{row[j]:.4f}' for j in dist_idx))

    # ---- 保存完整结果到 result1.xlsx ----
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output')
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, 'result1.xlsx')

    wb = openpyxl.Workbook()
    for sheet, hist in [('温度', T_hist), ('水分浓度', C_hist)]:
        ws = wb.create_sheet(sheet)
        ws.cell(1, 1, '时间\\到药材中心的距离')
        for j in range(N + 1):
            ws.cell(1, j + 2, round(r[j] * 100, 1))   # 列头：0,0.1,...,2.0 cm
        for t in range(int(t_end) + 1):
            ws.cell(t + 2, 1, t)                      # 行头：时间 s
            for j in range(N + 1):
                ws.cell(t + 2, j + 2, round(float(hist[t, j]), 4))
    # 删除默认 sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    wb.save(out)
    print(f'\n已保存完整结果到: {out}')

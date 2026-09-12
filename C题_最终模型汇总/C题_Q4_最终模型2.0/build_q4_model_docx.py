import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
Q3_DIR = ROOT / "C题_Q3_最终模型"
sys.path.insert(0, str(Q3_DIR))

import build_q3_model_docx as renderer  # noqa: E402


SOURCE = HERE / "第四问最终模型.md"
OUTPUT = HERE / "第四问最终模型_公式版.docx"

TABLE_CAPTIONS = {
    "4. 符号说明": "表1  第四问主要符号说明",
    "10.1 不同电价信息口径的结果": "表2  不同电价信息口径的全年结果",
    "10.2 电价预测有效性": "表3  因果电价预测精度",
    "10.3 价格响应结果": "表4  储能价格响应统计",
    "10.4 尾部风险": "表5  紧急购电尾部风险",
    "11.1 问题4-2计划购电量": "表6  问题4-2指定日期计划购电结果",
    "11.2 问题4-3基准计划购电量": "表7  问题4-3指定日期基准计划结果",
    "11.3 问题4-3最终调整购电量": "表8  问题4-3指定日期最终调整结果",
    "11.4 问题4-2储能运行": "表9  问题4-2指定日期储能运行结果",
    "11.5 问题4-3储能运行": "表10  问题4-3指定日期储能运行结果",
    "11.6 紧急购电结果": "表11  指定日期紧急购电结果",
}

FIGURE_MARKERS = {
    "[[FIGURE_PRICE]]": ("__Q4_FIGURE_PRICE__", "图1  指定日期动态电价曲线"),
    "[[FIGURE_RESPONSE]]": (
        "__Q4_FIGURE_RESPONSE__",
        "图2  动态电价优化价值与储能价格响应",
    ),
}


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    renderer.paper.configure_document(document)
    document.core_properties.title = "问题四最终模型（公式版）"
    document.core_properties.subject = "波动电价下的问题二与问题三重计算"
    document.core_properties.author = "数学建模论文组"

    equation_number = 0
    current_heading = ""
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            renderer.add_title(document, line[2:].strip())
            current_heading = line[2:].strip()
            index += 1
            continue
        if line.startswith("#### "):
            current_heading = line[5:].strip()
            renderer.paper.add_heading(document, current_heading, 3)
            index += 1
            continue
        if line.startswith("### "):
            current_heading = line[4:].strip()
            renderer.paper.add_heading(document, current_heading, 2)
            index += 1
            continue
        if line.startswith("## "):
            current_heading = line[3:].strip()
            renderer.paper.add_heading(document, current_heading, 1)
            index += 1
            continue
        if line.startswith("$$"):
            formula_parts = [line[2:]]
            while not formula_parts[-1].endswith("$$"):
                index += 1
                formula_parts.append(lines[index].strip())
            formula_parts[-1] = formula_parts[-1][:-2]
            equation_number += 1
            renderer.paper.add_equation(
                document,
                " ".join(formula_parts).strip(),
                equation_number,
            )
            index += 1
            continue
        if line in FIGURE_MARKERS:
            marker, caption = FIGURE_MARKERS[line]
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(marker)
            renderer.paper.add_caption(document, caption, kind="figure")
            index += 1
            continue
        if (
            line.startswith("|")
            and index + 1 < len(lines)
            and renderer.is_table_separator(lines[index + 1])
        ):
            headers = renderer.split_table_row(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(renderer.split_table_row(lines[index]))
                index += 1
            renderer.add_markdown_table(
                document,
                headers,
                rows,
                TABLE_CAPTIONS.get(current_heading, f"表  {current_heading}"),
                symbol_table=current_heading == "4. 符号说明",
            )
            continue
        numbered = renderer.re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            renderer.add_numbered_item(document, numbered.group(1), numbered.group(2))
        else:
            renderer.add_mixed_paragraph(document, line)
        index += 1

    document.save(OUTPUT)
    return OUTPUT, equation_number


if __name__ == "__main__":
    output, equations = build_document()
    print(f"saved={output}")
    print(f"equations={equations}")

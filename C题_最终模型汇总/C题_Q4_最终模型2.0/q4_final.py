from __future__ import annotations

from copy import copy
from datetime import time
from pathlib import Path
from shutil import copyfile

import numpy as np
import openpyxl
import pulp


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
Q2_DIR = ROOT / "C题_Q2_最终模型"
Q3_DIR = ROOT / "C题_Q3_最终模型"

import sys

sys.path.insert(0, str(Q2_DIR))
sys.path.insert(0, str(Q3_DIR))
import q2_final  # noqa: E402
import q3_final  # noqa: E402


ATTACHMENT4 = ROOT / "C题" / "附件" / "附件4.xlsx"
TEMPLATE42 = ROOT / "C题" / "附件" / "附件5" / "result4-2.xlsx"
TEMPLATE43 = ROOT / "C题" / "附件" / "附件5" / "result4-3.xlsx"
OUTPUT42 = HERE / "result4-2.xlsx"
OUTPUT43 = HERE / "result4-3.xlsx"
METRICS = HERE / "q4_metrics.txt"
SUMMARY_XLSX = HERE / "第四问结果汇总.xlsx"

N = q2_final.N
REPORT_START = q2_final.REPORT_START
E_INITIAL = q2_final.E_INITIAL
E_MIN = q2_final.E_MIN
E_MAX = q2_final.E_MAX
ETA_C = q2_final.ETA_C
ETA_D = q2_final.ETA_D
Q_MAX = q2_final.Q_MAX
TOL = q2_final.TOL
EXPORT_TOL = q2_final.EXPORT_TOL
SOLVER = pulp.PULP_CBC_CMD(msg=False)
PRICE_SEASON_LAG = 7
Q42_PRICE_WEIGHT = 1.0
Q43_PRICE_WEIGHT = 0.5
TERMINAL_MODE = q3_final.FINAL_TERMINAL_MODE
TERMINAL_BAND = q3_final.FINAL_TERMINAL_BAND


def load_dynamic_prices(dates):
    workbook = openpyxl.load_workbook(ATTACHMENT4, data_only=True, read_only=True)
    sheet = workbook.active
    price_dates = []
    prices = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        price_dates.append(row[0])
        prices.append(q2_final.rotate_to_natural_day(float(value) for value in row[1 : N + 1]))
    workbook.close()
    prices = np.asarray(prices)
    if prices.shape != (len(dates), N):
        raise ValueError(f"附件4电价维度异常: {prices.shape}")
    if [str(value)[:10] for value in price_dates] != [str(value)[:10] for value in dates]:
        raise ValueError("附件4与附件2日期不一致")
    if np.any(prices <= 0):
        raise ValueError("附件4存在非正电价")
    return prices


def causal_price_horizon(
    prices,
    prior_prices,
    day,
    release_slot,
    slots,
    seasonal_weight=1.0,
):
    """Forecast prices using only lag-7 profiles and prices revealed by release time."""
    forecast = np.zeros(slots)
    for offset in range(slots):
        absolute_slot = day * N + release_slot + offset
        target_day, target_slot = divmod(absolute_slot, N)
        source_day = target_day - PRICE_SEASON_LAG
        if 0 <= source_day < len(prices):
            forecast[offset] = prices[source_day, target_slot]
        else:
            forecast[offset] = prior_prices[target_slot]

    seasonal_weight = float(seasonal_weight)
    if not 0.0 <= seasonal_weight <= 1.0:
        raise ValueError("seasonal_weight必须位于[0,1]")

    # At 6:00/12:00/18:00, correct the weekly baseline by the mean error already
    # observed that day. At 0:00 one point is too noisy to estimate a level shift.
    if release_slot > 0 and day >= PRICE_SEASON_LAG:
        observed_error = (
            prices[day, : release_slot + 1]
            - prices[day - PRICE_SEASON_LAG, : release_slot + 1]
        )
        forecast += float(np.mean(observed_error))
    prior_horizon = np.asarray(
        [prior_prices[(release_slot + offset) % N] for offset in range(slots)]
    )
    forecast = seasonal_weight * forecast + (1.0 - seasonal_weight) * prior_horizon
    forecast[0] = prices[day, release_slot]
    return np.maximum(forecast, 1e-6)


def known_price_horizon(prices, prior_prices, day, release_slot, slots):
    """Perfect-price benchmark; use a causal seasonal proxy beyond the data boundary."""
    known = causal_price_horizon(prices, prior_prices, day, release_slot, slots)
    for offset in range(slots):
        absolute_slot = day * N + release_slot + offset
        target_day, target_slot = divmod(absolute_slot, N)
        if target_day < len(prices):
            known[offset] = prices[target_day, target_slot]
    return known


def optimization_price_horizon(
    price_mode,
    optimization_prices,
    settlement_prices,
    prior_prices,
    day,
    release_slot,
    slots,
    causal_weight=1.0,
):
    if price_mode == "causal":
        return causal_price_horizon(
            settlement_prices,
            prior_prices,
            day,
            release_slot,
            slots,
            seasonal_weight=causal_weight,
        )
    if price_mode == "perfect":
        return known_price_horizon(
            settlement_prices,
            prior_prices,
            day,
            release_slot,
            slots,
        )
    if price_mode == "matrix":
        horizon = np.zeros(slots)
        for offset in range(slots):
            absolute_slot = day * N + release_slot + offset
            target_day, target_slot = divmod(absolute_slot, N)
            if target_day < len(optimization_prices):
                horizon[offset] = optimization_prices[target_day, target_slot]
            else:
                horizon[offset] = optimization_prices[-1, target_slot]
        return horizon
    raise ValueError(f"未知电价信息模式: {price_mode}")


def solve_q42_horizon(
    net_targets,
    horizon_prices,
    e_start,
    terminal_mode=TERMINAL_MODE,
    terminal_band=TERMINAL_BAND,
):
    days = len(net_targets)
    slots = days * N
    horizon_prices = np.asarray(horizon_prices, dtype=float)
    if horizon_prices.shape != (slots,):
        raise ValueError("问题4-2的预测域电价长度不匹配")

    model = pulp.LpProblem("Q4_dynamic_price_q2", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("buy", range(slots), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(slots), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(slots), 0, Q_MAX)
    surplus = pulp.LpVariable.dicts("surplus", range(slots), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(slots + 1), E_MIN, E_MAX)
    model += energy[0] == float(e_start)
    if terminal_mode == "fixed":
        model += energy[slots] == E_INITIAL
    elif terminal_mode == "band":
        model += energy[slots] >= E_INITIAL - float(terminal_band)
        model += energy[slots] <= E_INITIAL + float(terminal_band)
    else:
        raise ValueError(f"问题4-2不支持的终端SOC模式: {terminal_mode}")

    for day in range(days):
        for t in range(N):
            k = day * N + t
            model += buy[k] + discharge[k] == float(net_targets[day][t]) + charge[k] + surplus[k]
            model += energy[k + 1] == energy[k] + ETA_C * charge[k] - discharge[k] / ETA_D

    cost = pulp.lpSum(float(horizon_prices[k]) * buy[k] for k in range(slots))
    model.setObjective(cost)
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"问题4-2第一阶段求解失败: {pulp.LpStatus[status]}")
    optimum = pulp.value(cost)
    model += cost <= optimum + 1e-3
    model.setObjective(
        pulp.lpSum(charge[k] + discharge[k] + surplus[k] for k in range(slots))
    )
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"问题4-2第二阶段求解失败: {pulp.LpStatus[status]}")

    return {
        "buy": np.asarray([pulp.value(buy[t]) for t in range(N)]),
        "charge": np.asarray([pulp.value(charge[t]) for t in range(N)]),
        "discharge": np.asarray([pulp.value(discharge[t]) for t in range(N)]),
        "surplus": np.asarray([pulp.value(surplus[t]) for t in range(N)]),
        "energy": np.asarray([pulp.value(energy[t]) for t in range(N + 1)]),
    }


def simulate_q42(
    optimization_prices,
    settlement_prices,
    prior_prices,
    prior_load,
    prior_pv,
    dates,
    loads,
    pvs,
    label,
    price_mode="matrix",
    terminal_mode=TERMINAL_MODE,
    terminal_band=TERMINAL_BAND,
    causal_price_weight=Q42_PRICE_WEIGHT,
):
    residuals = []
    results = []
    e_start = E_INITIAL
    max_balance = 0.0
    max_storage = 0.0
    min_energy = E_MAX
    max_energy = E_MIN

    for day in range(len(dates)):
        load0, _ = q2_final.point_forecast(loads, day, 0, prior_load)
        pv0, _ = q2_final.point_forecast(pvs, day, 0, prior_pv)
        margin = q2_final.residual_margin(residuals)
        load1, _ = q2_final.point_forecast(loads, day, 1, prior_load)
        pv1, _ = q2_final.point_forecast(pvs, day, 1, prior_pv)
        targets = [load0 - pv0 + margin, load1 - pv1 + margin]
        price_horizon = optimization_price_horizon(
            price_mode,
            optimization_prices,
            settlement_prices,
            prior_prices,
            day,
            0,
            2 * N,
            causal_weight=causal_price_weight,
        )
        plan = solve_q42_horizon(
            targets,
            price_horizon,
            e_start,
            terminal_mode=terminal_mode,
            terminal_band=terminal_band,
        )
        actual = q2_final.execute_causal_day(plan["buy"], loads[day], pvs[day], e_start)
        result = {
            "date": dates[day],
            "buy": plan["buy"],
            "charge": actual["charge"],
            "discharge": actual["discharge"],
            "emergency": actual["emergency"],
            "surplus": actual["surplus"],
            "energy": actual["energy"],
        }
        balance = (
            result["buy"]
            + pvs[day]
            + result["discharge"]
            + result["emergency"]
            - loads[day]
            - result["charge"]
            - result["surplus"]
        )
        storage = (
            result["energy"][1:]
            - result["energy"][:-1]
            - ETA_C * result["charge"]
            + result["discharge"] / ETA_D
        )
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_storage = max(max_storage, float(np.max(np.abs(storage))))
        min_energy = min(min_energy, float(np.min(result["energy"])))
        max_energy = max(max_energy, float(np.max(result["energy"])))
        if np.any((result["charge"] > TOL) & (result["discharge"] > TOL)):
            raise AssertionError(f"{label}存在同时充放电")
        results.append(result)
        residuals.append((loads[day] - pvs[day]) - (load0 - pv0))
        e_start = float(result["energy"][-1])
        if (day + 1) % 60 == 0 or day + 1 == len(dates):
            print(f"{label}: {day + 1}/{len(dates)}", flush=True)

    for day, result in enumerate(results):
        price = settlement_prices[day]
        result["scheduled_cost"] = float(np.dot(price, result["buy"]))
        result["emergency_cost"] = float(np.dot(5.0 * price, result["emergency"]))
        result["total_cost"] = result["scheduled_cost"] + result["emergency_cost"]
    return results, {
        "balance": max_balance,
        "storage": max_storage,
        "min_energy": min_energy,
        "max_energy": max_energy,
    }


def summarize_q42(results):
    report = results[REPORT_START:]
    daily_emergency = np.asarray(
        [float(np.sum(result["emergency"])) for result in report]
    )
    daily_emergency_cost = np.asarray(
        [float(result["emergency_cost"]) for result in report]
    )
    emergency_slots = (
        np.concatenate([result["emergency"] for result in report])
        if report
        else np.array([])
    )
    if len(daily_emergency_cost):
        cost_var95 = float(np.quantile(daily_emergency_cost, 0.95))
        tail_costs = daily_emergency_cost[daily_emergency_cost >= cost_var95 - 1e-12]
        cost_cvar95 = float(np.mean(tail_costs))
    else:
        cost_var95 = np.nan
        cost_cvar95 = np.nan
    return {
        "days": len(report),
        "purchase_kwh": sum(float(np.sum(result["buy"])) for result in report),
        "emergency_kwh": sum(float(np.sum(result["emergency"])) for result in report),
        "surplus_kwh": sum(float(np.sum(result["surplus"])) for result in report),
        "scheduled_cost": sum(result["scheduled_cost"] for result in report),
        "emergency_cost": sum(result["emergency_cost"] for result in report),
        "total_cost": sum(result["total_cost"] for result in report),
        "emergency_days": sum(float(np.sum(result["emergency"])) > 0.01 for result in report),
        "daily_emergency_p95_kwh": (
            float(np.quantile(daily_emergency, 0.95)) if len(daily_emergency) else np.nan
        ),
        "max_interval_emergency_kwh": (
            float(np.max(emergency_slots)) if len(emergency_slots) else np.nan
        ),
        "emergency_duration_hours": (
            float(np.sum(emergency_slots > EXPORT_TOL) * q3_final.DT)
            if len(emergency_slots)
            else 0.0
        ),
        "daily_emergency_cost_var95": cost_var95,
        "daily_emergency_cost_cvar95": cost_cvar95,
        "end_energy": float(report[-1]["energy"][-1]),
    }


def simulate_q43(
    optimization_prices,
    settlement_prices,
    prior_prices,
    prior_load,
    dates,
    loads,
    pvs,
    pv_forecasts,
    net_load_margins,
    label,
    price_mode="matrix",
    terminal_mode=TERMINAL_MODE,
    terminal_band=TERMINAL_BAND,
    causal_price_weight=Q43_PRICE_WEIGHT,
):
    results = []
    e_start = E_INITIAL
    max_balance = 0.0
    max_storage = 0.0
    min_energy = E_MAX
    max_energy = E_MIN

    for day in range(len(dates)):
        daily_plan = None
        adjusted = np.zeros(N)
        charge = np.zeros(N)
        discharge = np.zeros(N)
        emergency = np.zeros(N)
        spill = np.zeros(N)
        energy = np.zeros(N + 1)
        energy[0] = e_start
        current = float(e_start)
        for release_position, release_slot in enumerate(q3_final.RELEASE_SLOTS):
            release_index = q3_final.RELEASE_SLOTS.index(release_slot)
            load_horizon = q3_final.load_forecast_horizon(
                loads, day, release_slot, prior_load
            )
            load_horizon = np.maximum(
                load_horizon + net_load_margins[day, release_index],
                0.0,
            )
            pv_horizon = pv_forecasts[day, release_index].copy()
            price_horizon = optimization_price_horizon(
                price_mode,
                optimization_prices,
                settlement_prices,
                prior_prices,
                day,
                release_slot,
                N,
                causal_weight=causal_price_weight,
            )
            base = None if release_slot == 0 else daily_plan[release_slot:]
            schedule = q3_final.solve_schedule(
                load_horizon,
                pv_horizon,
                price_horizon,
                current,
                base,
                terminal_mode=terminal_mode,
                terminal_band=terminal_band,
            )
            if release_slot == 0:
                daily_plan = schedule["buy"].copy()

            next_release = (
                q3_final.RELEASE_SLOTS[release_position + 1]
                if release_position + 1 < len(q3_final.RELEASE_SLOTS)
                else N
            )
            block_length = next_release - release_slot
            stop = release_slot + block_length
            committed = schedule["buy"][:block_length]
            adjusted[release_slot:stop] = committed
            actual = q3_final.execute_block(
                committed,
                loads[day, release_slot:stop],
                pvs[day, release_slot:stop],
                current,
            )
            charge[release_slot:stop] = actual["charge"]
            discharge[release_slot:stop] = actual["discharge"]
            emergency[release_slot:stop] = actual["emergency"]
            spill[release_slot:stop] = actual["spill"]
            energy[release_slot : stop + 1] = actual["energy"]
            current = float(actual["energy"][-1])

        balance = adjusted + pvs[day] + discharge + emergency - loads[day] - charge - spill
        storage = energy[1:] - energy[:-1] - ETA_C * charge + discharge / ETA_D
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_storage = max(max_storage, float(np.max(np.abs(storage))))
        min_energy = min(min_energy, float(np.min(energy)))
        max_energy = max(max_energy, float(np.max(energy)))
        if np.any((charge > TOL) & (discharge > TOL)):
            raise AssertionError(f"{label}存在同时充放电")

        costs = q3_final.settlement(
            settlement_prices[day],
            daily_plan,
            adjusted,
            emergency,
        )
        results.append(
            {
                "date": dates[day],
                "plan": daily_plan,
                "adjusted": adjusted,
                "charge": charge,
                "discharge": discharge,
                "emergency": emergency,
                "spill": spill,
                "energy": energy,
                **costs,
            }
        )
        e_start = current
        if (day + 1) % 60 == 0 or day + 1 == len(dates):
            print(f"{label}: {day + 1}/{len(dates)}", flush=True)

    return results, {
        "balance": max_balance,
        "storage": max_storage,
        "min_energy": min_energy,
        "max_energy": max_energy,
    }


def write_storage_emergency(workbook, dates, results, emergency_key="emergency"):
    sheet = workbook["充放电量"]
    styles = [[copy(sheet.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    sheet.delete_rows(2, sheet.max_row)
    labels = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]
    row_index = 2
    for date_value, result in zip(dates, results):
        for block in range(6):
            for column in range(1, 7):
                sheet.cell(row_index, column)._style = copy(styles[block][column - 1])
            start, stop = block * 24, (block + 1) * 24
            sheet.cell(row_index, 1, date_value if block == 0 else None)
            sheet.cell(row_index, 2, labels[block])
            sheet.cell(row_index, 3, round(float(np.sum(result["charge"][start:stop])), 4))
            sheet.cell(row_index, 4, round(float(np.sum(result["discharge"][start:stop])), 4))
            if block == 0:
                sheet.cell(row_index, 5, time(0, 0))
                sheet.cell(row_index, 6, round(float(result["energy"][0]), 4))
            elif block == 1:
                sheet.cell(row_index, 5, "24:00")
                sheet.cell(row_index, 6, round(float(result["energy"][-1]), 4))
            row_index += 1

    sheet = workbook["紧急购电量"]
    styles = [[copy(sheet.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    sheet.delete_rows(2, sheet.max_row)
    row_index = 2
    for date_value, result in zip(dates, results):
        emergency = result[emergency_key]
        events = []
        t = 0
        while t < N:
            if emergency[t] < EXPORT_TOL:
                t += 1
                continue
            start = t
            while t < N and emergency[t] >= EXPORT_TOL:
                t += 1
            events.append((start, t))
        for event_index, (start, stop) in enumerate(events):
            style = styles[min(event_index, len(styles) - 1)]
            for column in range(1, 4):
                sheet.cell(row_index, column)._style = copy(style[column - 1])
            sheet.cell(row_index, 1, date_value if event_index == 0 else None)
            sheet.cell(row_index, 2, q2_final.interval_label(start, stop))
            sheet.cell(row_index, 3, round(float(np.sum(emergency[start:stop])), 4))
            row_index += 1


def write_q42(dates, prices, results):
    report = results[REPORT_START:]
    report_dates = dates[REPORT_START:]
    report_prices = prices[REPORT_START:]
    copyfile(TEMPLATE42, OUTPUT42)
    workbook = openpyxl.load_workbook(OUTPUT42)
    sheet = workbook["计划购电量"]
    for row_index, (price, result) in enumerate(zip(report_prices, report), start=2):
        for column, value in enumerate(q2_final.source_order(result["buy"]), start=2):
            sheet.cell(row_index, column, round(float(value), 4))
        sheet.cell(row_index, 146, round(float(np.sum(result["buy"])), 4))
        sheet.cell(row_index, 147, round(float(np.dot(price, result["buy"])), 4))
    write_storage_emergency(workbook, report_dates, report)
    workbook.save(OUTPUT42)


def write_q43(dates, results):
    report = results[REPORT_START:]
    report_dates = dates[REPORT_START:]
    copyfile(TEMPLATE43, OUTPUT43)
    workbook = openpyxl.load_workbook(OUTPUT43)
    for sheet_name, key, cost_key in (
        ("计划购电量", "plan", "plan_cost"),
        ("调整购电量", "adjusted", "scheduled_cost"),
    ):
        sheet = workbook[sheet_name]
        for row_index, result in enumerate(report, start=2):
            for column, value in enumerate(q2_final.source_order(result[key]), start=2):
                sheet.cell(row_index, column, round(float(value), 4))
            sheet.cell(row_index, 146, round(float(np.sum(result[key])), 4))
            sheet.cell(row_index, 147, round(float(result[cost_key]), 4))
    write_storage_emergency(workbook, report_dates, report)
    workbook.save(OUTPUT43)


def price_response(prices, results):
    report_prices = prices[REPORT_START:].reshape(-1)
    charge = np.concatenate([result["charge"] for result in results[REPORT_START:]])
    discharge = np.concatenate([result["discharge"] for result in results[REPORT_START:]])
    q25, q75 = np.quantile(report_prices, [0.25, 0.75])
    average_charge_price = float(np.dot(report_prices, charge) / max(np.sum(charge), 1e-12))
    average_discharge_price = float(
        np.dot(report_prices, discharge) / max(np.sum(discharge), 1e-12)
    )
    low_charge_share = float(np.sum(charge[report_prices <= q25]) / max(np.sum(charge), 1e-12))
    high_discharge_share = float(
        np.sum(discharge[report_prices >= q75]) / max(np.sum(discharge), 1e-12)
    )
    net_output = discharge - charge
    correlation = float(np.corrcoef(report_prices, net_output)[0, 1])
    return {
        "average_charge_price": average_charge_price,
        "average_discharge_price": average_discharge_price,
        "low_price_charge_share": low_charge_share,
        "high_price_discharge_share": high_discharge_share,
        "price_net_output_correlation": correlation,
        "q25": float(q25),
        "q75": float(q75),
    }


def price_forecast_quality(prices, prior_prices):
    q42_forecasts = []
    q42_actuals = []
    q43_forecasts = []
    q43_actuals = []
    for day in range(REPORT_START, len(prices)):
        forecast = causal_price_horizon(
            prices,
            prior_prices,
            day,
            0,
            N,
            seasonal_weight=Q42_PRICE_WEIGHT,
        )
        q42_forecasts.append(forecast)
        q42_actuals.append(prices[day])
        for release_slot in q3_final.RELEASE_SLOTS:
            stop = min(release_slot + q3_final.BLOCK, N)
            block_length = stop - release_slot
            forecast = causal_price_horizon(
                prices,
                prior_prices,
                day,
                release_slot,
                N,
                seasonal_weight=Q43_PRICE_WEIGHT,
            )[:block_length]
            q43_forecasts.append(forecast)
            q43_actuals.append(prices[day, release_slot:stop])

    def metrics(forecasts, actuals):
        forecast = np.concatenate(forecasts)
        actual = np.concatenate(actuals)
        return {
            "mae": float(np.mean(np.abs(forecast - actual))),
            "rmse": float(np.sqrt(np.mean((forecast - actual) ** 2))),
            "correlation": float(np.corrcoef(forecast, actual)[0, 1]),
        }

    q42 = metrics(q42_forecasts, q42_actuals)
    q43 = metrics(q43_forecasts, q43_actuals)
    return {
        "q42_mae": q42["mae"],
        "q42_rmse": q42["rmse"],
        "q42_correlation": q42["correlation"],
        "q43_block_mae": q43["mae"],
        "q43_block_rmse": q43["rmse"],
        "q43_block_correlation": q43["correlation"],
    }


def no_storage_physical_reference(prices, loads, pvs):
    """Physical reference only: actual net demand bought at the spot price."""
    net_demand = np.maximum(loads[REPORT_START:] - pvs[REPORT_START:], 0.0)
    report_prices = prices[REPORT_START:]
    return {
        "purchase_kwh": float(np.sum(net_demand)),
        "total_cost": float(np.sum(report_prices * net_demand)),
    }


def create_summary_workbook(
    price_stats,
    summaries,
    responses,
    forecast_quality,
    no_storage,
):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "模型结果"
    sheet.append(
        [
            "模型",
            "正常购电费/万元",
            "紧急购电费/万元",
            "总费用/万元",
            "紧急购电量/kWh",
            "紧急购电天数",
        ]
    )
    for name in (
        "q42_fixed",
        "q42_causal",
        "q42_perfect",
        "q43_fixed",
        "q43_causal",
        "q43_perfect",
    ):
        values = summaries[name]
        sheet.append(
            [
                name,
                values["scheduled_cost"] / 10000,
                values["emergency_cost"] / 10000,
                values["total_cost"] / 10000,
                values["emergency_kwh"],
                values["emergency_days"],
            ]
        )

    sheet2 = workbook.create_sheet("价格响应")
    sheet2.append(
        [
            "模型",
            "加权平均充电电价",
            "加权平均放电电价",
            "低价四分位充电占比",
            "高价四分位放电占比",
            "电价-净放电相关系数",
        ]
    )
    for name in ("q42", "q43"):
        values = responses[name]
        sheet2.append(
            [
                name,
                values["average_charge_price"],
                values["average_discharge_price"],
                values["low_price_charge_share"],
                values["high_price_discharge_share"],
                values["price_net_output_correlation"],
            ]
        )

    sheet3 = workbook.create_sheet("电价特征")
    sheet3.append(["指标", "数值"])
    for key, value in price_stats.items():
        sheet3.append([key, value])

    sheet4 = workbook.create_sheet("电价预测检验")
    sheet4.append(["指标", "数值"])
    for key, value in forecast_quality.items():
        sheet4.append([key, value])

    sheet5 = workbook.create_sheet("物理参照")
    sheet5.append(["指标", "数值"])
    for key, value in no_storage.items():
        sheet5.append([key, value])

    thin = openpyxl.styles.Side(style="thin", color="000000")
    medium = openpyxl.styles.Side(style="medium", color="000000")
    for current in workbook.worksheets:
        for cell in current[1]:
            cell.font = openpyxl.styles.Font(name="宋体", size=10.5, bold=True)
            cell.alignment = openpyxl.styles.Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = openpyxl.styles.Border(top=medium, bottom=thin)
        for row in current.iter_rows(min_row=2):
            for cell in row:
                cell.font = openpyxl.styles.Font(name="Times New Roman", size=10)
                cell.alignment = openpyxl.styles.Alignment(
                    horizontal="center", vertical="center"
                )
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for cell in current[current.max_row]:
            cell.border = openpyxl.styles.Border(bottom=medium)
        for column in range(1, current.max_column + 1):
            current.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 23
        current.freeze_panes = "A2"
    workbook.save(SUMMARY_XLSX)


def main():
    fixed_prices, prior_load, prior_pv, dates, loads, pvs = q2_final.load_inputs()
    dynamic_prices = load_dynamic_prices(dates)
    fixed_price_matrix = np.tile(fixed_prices, (len(dates), 1))

    q42_causal, q42_causal_checks = simulate_q42(
        dynamic_prices,
        dynamic_prices,
        fixed_prices,
        prior_load,
        prior_pv,
        dates,
        loads,
        pvs,
        "q4-2_causal_price",
        price_mode="causal",
    )
    q42_perfect, q42_perfect_checks = simulate_q42(
        dynamic_prices,
        dynamic_prices,
        fixed_prices,
        prior_load,
        prior_pv,
        dates,
        loads,
        pvs,
        "q4-2_perfect_price",
        price_mode="perfect",
    )
    q42_fixed, q42_fixed_checks = simulate_q42(
        fixed_price_matrix,
        dynamic_prices,
        fixed_prices,
        prior_load,
        prior_pv,
        dates,
        loads,
        pvs,
        "q4-2_fixed_benchmark",
        price_mode="matrix",
    )

    pv_forecasts, _ = q3_final.load_pv_forecasts(dates, pvs)
    net_load_margins = q3_final.build_net_load_margins(
        loads,
        pvs,
        pv_forecasts,
        prior_load,
    )
    q43_causal, q43_causal_checks = simulate_q43(
        dynamic_prices,
        dynamic_prices,
        fixed_prices,
        prior_load,
        dates,
        loads,
        pvs,
        pv_forecasts,
        net_load_margins,
        "q4-3_causal_price",
        price_mode="causal",
    )
    q43_perfect, q43_perfect_checks = simulate_q43(
        dynamic_prices,
        dynamic_prices,
        fixed_prices,
        prior_load,
        dates,
        loads,
        pvs,
        pv_forecasts,
        net_load_margins,
        "q4-3_perfect_price",
        price_mode="perfect",
    )
    q43_fixed, q43_fixed_checks = simulate_q43(
        fixed_price_matrix,
        dynamic_prices,
        fixed_prices,
        prior_load,
        dates,
        loads,
        pvs,
        pv_forecasts,
        net_load_margins,
        "q4-3_fixed_benchmark",
        price_mode="matrix",
    )

    write_q42(dates, dynamic_prices, q42_causal)
    write_q43(dates, q43_causal)

    summaries = {
        "q42_causal": summarize_q42(q42_causal),
        "q42_perfect": summarize_q42(q42_perfect),
        "q42_fixed": summarize_q42(q42_fixed),
        "q43_causal": q3_final.summarize(q43_causal),
        "q43_perfect": q3_final.summarize(q43_perfect),
        "q43_fixed": q3_final.summarize(q43_fixed),
    }
    responses = {
        "q42": price_response(dynamic_prices, q42_causal),
        "q43": price_response(dynamic_prices, q43_causal),
    }
    report_prices = dynamic_prices[REPORT_START:]
    daily_min = np.min(report_prices, axis=1)
    daily_max = np.max(report_prices, axis=1)
    price_stats = {
        "minimum_price": float(np.min(report_prices)),
        "maximum_price": float(np.max(report_prices)),
        "mean_price": float(np.mean(report_prices)),
        "standard_deviation": float(np.std(report_prices)),
        "coefficient_of_variation": float(np.std(report_prices) / np.mean(report_prices)),
        "mean_daily_range": float(np.mean(daily_max - daily_min)),
        "arbitrage_feasible_days": int(np.sum(daily_max > daily_min / (ETA_C * ETA_D))),
        "report_days": int(len(report_prices)),
    }
    forecast_quality = price_forecast_quality(dynamic_prices, fixed_prices)
    no_storage = no_storage_physical_reference(dynamic_prices, loads, pvs)
    create_summary_workbook(
        price_stats,
        summaries,
        responses,
        forecast_quality,
        no_storage,
    )

    checks = {
        "q42_causal": q42_causal_checks,
        "q42_perfect": q42_perfect_checks,
        "q42_fixed": q42_fixed_checks,
        "q43_causal": q43_causal_checks,
        "q43_perfect": q43_perfect_checks,
        "q43_fixed": q43_fixed_checks,
    }
    lines = ["Q4 FINAL MODEL METRICS"]
    for name, values in summaries.items():
        lines.append(f"[{name}]")
        lines.extend(f"{key}={value}" for key, value in values.items())
        lines.extend(f"{key}={value}" for key, value in checks[name].items())
    lines.append("[price_statistics]")
    lines.extend(f"{key}={value}" for key, value in price_stats.items())
    lines.append("[price_forecast_quality]")
    lines.extend(f"{key}={value}" for key, value in forecast_quality.items())
    lines.append("[no_storage_physical_reference]")
    lines.extend(f"{key}={value}" for key, value in no_storage.items())
    for name, values in responses.items():
        lines.append(f"[price_response_{name}]")
        lines.extend(f"{key}={value}" for key, value in values.items())
    lines.append("[optimization_value]")
    lines.append(
        "q42_causal_adaptation_saving="
        f"{summaries['q42_fixed']['total_cost'] - summaries['q42_causal']['total_cost']}"
    )
    lines.append(
        "q43_causal_adaptation_saving="
        f"{summaries['q43_fixed']['total_cost'] - summaries['q43_causal']['total_cost']}"
    )
    lines.append(
        "q42_perfect_information_gap="
        f"{summaries['q42_causal']['total_cost'] - summaries['q42_perfect']['total_cost']}"
    )
    lines.append(
        "q43_perfect_information_gap="
        f"{summaries['q43_causal']['total_cost'] - summaries['q43_perfect']['total_cost']}"
    )
    lines.append(f"output42={OUTPUT42}")
    lines.append(f"output43={OUTPUT43}")
    METRICS.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()

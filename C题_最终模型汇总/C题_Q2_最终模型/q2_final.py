from __future__ import annotations

from copy import copy
from datetime import datetime, time, timedelta
from pathlib import Path
from shutil import copyfile

import numpy as np
import openpyxl
import pulp


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT1 = ROOT / "C题" / "附件" / "附件1.xlsx"
ATTACHMENT2 = ROOT / "C题" / "附件" / "附件2.xlsx"
TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result2.xlsx"
OUTPUT = Path(__file__).resolve().parent / "result2_optimized.xlsx"
SUMMARY = Path(__file__).resolve().parent / "q2_optimized_metrics.txt"

DT = 1.0 / 6.0
N = 144
ETA_C = 0.9
ETA_D = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
Q_MAX = 5000.0 * DT
REPORT_START = 31
RISK_QUANTILE = 0.80
RESIDUAL_WINDOW = 28
ALPHA_GRID = np.linspace(0.0, 1.0, 11)
SOLVER = pulp.PULP_CBC_CMD(msg=False)
TOL = 1e-5
EXPORT_TOL = 5e-5


def rotate_to_natural_day(values):
    """Convert template order (00:10,...,24:00,00:10+1) to 00:00,...,24:00."""
    values = list(values)
    if len(values) != N:
        raise ValueError(f"Expected {N} intervals, got {len(values)}")
    return np.asarray([values[-1], *values[:-1]], dtype=float)


def source_order(values):
    """Convert natural-day order back to the official workbook column order."""
    values = list(values)
    return [*values[1:], values[0]]


def load_inputs():
    wb1 = openpyxl.load_workbook(ATTACHMENT1, data_only=True, read_only=True)
    rows = list(wb1["Sheet1"].iter_rows(min_row=2, values_only=True))
    wb1.close()
    if len(rows) != N:
        raise ValueError("附件1的时段数不是144")
    prices = rotate_to_natural_day(float(row[1]) for row in rows)
    prior_load = rotate_to_natural_day(float(row[2]) * DT for row in rows)
    prior_pv = rotate_to_natural_day(float(row[3]) * DT for row in rows)

    wb2 = openpyxl.load_workbook(ATTACHMENT2, data_only=True, read_only=True)
    load_sheet, pv_sheet = wb2.worksheets[:2]

    def read_sheet(sheet):
        data = []
        dates = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            dates.append(row[0])
            data.append(rotate_to_natural_day(float(v) * DT for v in row[1 : N + 1]))
        return dates, np.asarray(data)

    dates, loads = read_sheet(load_sheet)
    pv_dates, pvs = read_sheet(pv_sheet)
    wb2.close()
    if len(dates) != 365 or loads.shape != (365, N) or pvs.shape != (365, N):
        raise ValueError(f"附件2维度异常: load={loads.shape}, pv={pvs.shape}")
    if [str(x) for x in dates] != [str(x) for x in pv_dates]:
        raise ValueError("负载与光伏日期不一致")
    return prices, prior_load, prior_pv, dates, loads, pvs


def select_alpha(data, day, horizon, prior):
    """Choose the lag blend only from observations available at day 00:00."""
    target = day + horizon
    max_validation_day = day - 1
    start = max(7, max_validation_day - 55)
    candidates = []
    for alpha in ALPHA_GRID:
        errors = []
        for k in range(start, max_validation_day + 1):
            if horizon == 0:
                pred = alpha * data[k - 1] + (1.0 - alpha) * data[k - 7]
            else:
                # Reconstruct a two-day-ahead forecast made at k-1.
                if k < 7:
                    continue
                pred = alpha * data[k - 2] + (1.0 - alpha) * data[k - 7]
            errors.append(np.mean(np.abs(data[k] - pred)))
        candidates.append(np.mean(errors) if errors else np.inf)
    if not np.isfinite(min(candidates)):
        return 0.0
    return float(ALPHA_GRID[int(np.argmin(candidates))])


def point_forecast(data, day, horizon, prior):
    target = day + horizon
    if day < 7:
        return prior.copy(), 0.0
    alpha = select_alpha(data, day, horizon, prior)
    if horizon == 0:
        forecast = alpha * data[day - 1] + (1.0 - alpha) * data[day - 7]
    else:
        forecast = alpha * data[day - 1] + (1.0 - alpha) * data[day - 6]
    return np.maximum(forecast, 0.0), alpha


def residual_margin(residual_history):
    if not residual_history:
        return np.zeros(N)
    sample = np.asarray(residual_history[-RESIDUAL_WINDOW:])
    return np.quantile(sample, RISK_QUANTILE, axis=0)


def solve_horizon(net_targets, prices, e_start, terminal_target=E_INITIAL):
    horizon = len(net_targets)
    total_slots = horizon * N
    model = pulp.LpProblem("Q2_rolling_plan", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("buy", range(total_slots), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(total_slots), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(total_slots), 0, Q_MAX)
    surplus = pulp.LpVariable.dicts("surplus", range(total_slots), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(total_slots + 1), E_MIN, E_MAX)
    model += energy[0] == e_start
    model += energy[total_slots] == terminal_target

    for j in range(horizon):
        for t in range(N):
            k = j * N + t
            model += buy[k] + discharge[k] == float(net_targets[j][t]) + charge[k] + surplus[k]
            model += energy[k + 1] == energy[k] + ETA_C * charge[k] - discharge[k] / ETA_D

    cost = pulp.lpSum(prices[k % N] * buy[k] for k in range(total_slots))
    model.setObjective(cost)
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"滚动计划第一阶段求解失败: {pulp.LpStatus[status]}")
    optimum = pulp.value(cost)

    # CBC stores an objective of roughly 1e5 yuan per horizon; a 1e-3 yuan
    # lexicographic tolerance is far below reporting precision and numerically stable.
    model += cost <= optimum + 1e-3
    model.setObjective(
        pulp.lpSum(charge[k] + discharge[k] + surplus[k] for k in range(total_slots))
    )
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"滚动计划第二阶段求解失败: {pulp.LpStatus[status]}")

    first = range(N)
    return {
        "buy": np.array([pulp.value(buy[t]) for t in first]),
        "charge": np.array([pulp.value(charge[t]) for t in first]),
        "discharge": np.array([pulp.value(discharge[t]) for t in first]),
        "planned_surplus": np.array([pulp.value(surplus[t]) for t in first]),
        "energy": np.array([pulp.value(energy[t]) for t in range(N + 1)]),
    }


def execute_fixed_plan(plan, load, pv):
    """Causal settlement: the 00:00 purchase/storage schedule remains fixed."""
    gap = load + plan["charge"] - plan["buy"] - pv - plan["discharge"]
    emergency = np.maximum(gap, 0.0)
    surplus = np.maximum(-gap, 0.0)
    return emergency, surplus


def execute_causal_day(buy, load, pv, e_start):
    """Use only the current interval observation to balance forecast errors."""
    energy = float(e_start)
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    surplus = np.zeros(N)
    trajectory = np.zeros(N + 1)
    trajectory[0] = energy
    for t in range(N):
        net_deficit = load[t] - pv[t] - buy[t]
        if net_deficit >= 0.0:
            discharge[t] = min(net_deficit, Q_MAX, max((energy - E_MIN) * ETA_D, 0.0))
            emergency[t] = net_deficit - discharge[t]
        else:
            available = -net_deficit
            charge[t] = min(available, Q_MAX, max((E_MAX - energy) / ETA_C, 0.0))
            surplus[t] = available - charge[t]
        energy += ETA_C * charge[t] - discharge[t] / ETA_D
        trajectory[t + 1] = energy
    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "surplus": surplus,
        "energy": trajectory,
    }


def validate_day(plan, load, pv, emergency, surplus):
    balance = plan["buy"] + pv + plan["discharge"] + emergency - load - plan["charge"] - surplus
    storage = (
        plan["energy"][1:]
        - plan["energy"][:-1]
        - ETA_C * plan["charge"]
        + plan["discharge"] / ETA_D
    )
    simultaneous = np.flatnonzero((plan["charge"] > TOL) & (plan["discharge"] > TOL))
    checks = {
        "balance": float(np.max(np.abs(balance))),
        "storage": float(np.max(np.abs(storage))),
        "min_energy": float(np.min(plan["energy"])),
        "max_energy": float(np.max(plan["energy"])),
        "simultaneous": int(len(simultaneous)),
    }
    if checks["balance"] > 1e-3 or checks["storage"] > 1e-3:
        raise AssertionError(f"约束残差过大: {checks}")
    if checks["min_energy"] < E_MIN - 1e-3 or checks["max_energy"] > E_MAX + 1e-3:
        raise AssertionError(f"SOC越界: {checks}")
    if checks["simultaneous"]:
        raise AssertionError(f"存在同时充放电: {simultaneous.tolist()}")
    return checks


def copy_row_style(sheet, source_row, target_row, max_col):
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        src = sheet.cell(source_row, col)
        dst = sheet.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)


def interval_label(start, stop):
    def hhmm(minutes):
        minutes %= 1440
        return f"{minutes // 60:02d}:{minutes % 60:02d}"

    start_min = start * 10
    stop_min = stop * 10
    suffix = "+1" if stop_min >= 1440 else ""
    return f"{hhmm(start_min)}-{hhmm(stop_min)}{suffix}"


def write_output(dates, prices, results):
    copyfile(TEMPLATE, OUTPUT)
    wb = openpyxl.load_workbook(OUTPUT)

    ws = wb["计划购电量"]
    for i, result in enumerate(results):
        row = i + 2
        values = source_order(result["buy"])
        for t, value in enumerate(values, start=2):
            ws.cell(row, t, round(float(value), 4))
        ws.cell(row, 146, round(float(np.sum(result["buy"])), 4))
        ws.cell(row, 147, round(float(np.dot(prices, result["buy"])), 4))

    ws = wb["充放电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    ws.delete_rows(2, ws.max_row)
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for date, result in zip(dates, results):
        for block in range(6):
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(style_rows[block][col - 1])
            start, stop = block * 24, (block + 1) * 24
            ws.cell(row, 1, date if block == 0 else None)
            ws.cell(row, 2, labels[block])
            ws.cell(row, 3, round(float(np.sum(result["charge"][start:stop])), 4))
            ws.cell(row, 4, round(float(np.sum(result["discharge"][start:stop])), 4))
            if block == 0:
                ws.cell(row, 5, time(0, 0))
                ws.cell(row, 6, round(float(result["energy"][0]), 4))
            elif block == 1:
                ws.cell(row, 5, "24:00")
                ws.cell(row, 6, round(float(result["energy"][-1]), 4))
            row += 1

    ws = wb["紧急购电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    ws.delete_rows(2, ws.max_row)
    row = 2
    for date, result in zip(dates, results):
        emergency = result["emergency"]
        starts = []
        t = 0
        while t < N:
            if emergency[t] < EXPORT_TOL:
                t += 1
                continue
            start = t
            while t < N and emergency[t] >= EXPORT_TOL:
                t += 1
            starts.append((start, t))
        for event_index, (start, stop) in enumerate(starts):
            style = style_rows[min(event_index, len(style_rows) - 1)]
            for col in range(1, 4):
                ws.cell(row, col)._style = copy(style[col - 1])
            ws.cell(row, 1, date if event_index == 0 else None)
            ws.cell(row, 2, interval_label(start, stop))
            ws.cell(row, 3, round(float(np.sum(emergency[start:stop])), 4))
            row += 1

    wb.save(OUTPUT)


def metric_line(name, actual, pred):
    err = actual - pred
    mae = np.mean(np.abs(err))
    rmse = np.sqrt(np.mean(err**2))
    wmape = np.sum(np.abs(err)) / max(np.sum(np.abs(actual)), 1e-12)
    return f"{name}: MAE={mae:.4f} kWh, RMSE={rmse:.4f} kWh, WMAPE={wmape:.2%}"


def main():
    prices, prior_load, prior_pv, dates, loads, pvs = load_inputs()
    residuals = []
    all_results = []
    point_loads = []
    point_pvs = []
    persistence_loads = []
    persistence_pvs = []
    max_checks = {"balance": 0.0, "storage": 0.0, "min_energy": E_MAX, "max_energy": E_MIN}
    e_start = E_INITIAL

    for day in range(len(dates)):
        load0, alpha_l0 = point_forecast(loads, day, 0, prior_load)
        pv0, alpha_g0 = point_forecast(pvs, day, 0, prior_pv)
        net0 = load0 - pv0
        margin = residual_margin(residuals)
        targets = [net0 + margin]

        # Keep a full 48-hour look-ahead on 31 December as well. The auxiliary
        # 1 January forecast uses only already observed lagged data and avoids an
        # artificial year-end discharge caused by truncating the horizon.
        load1, _ = point_forecast(loads, day, 1, prior_load)
        pv1, _ = point_forecast(pvs, day, 1, prior_pv)
        targets.append(load1 - pv1 + margin)

        plan = solve_horizon(targets, prices, e_start)
        fixed_emergency, fixed_surplus = execute_fixed_plan(plan, loads[day], pvs[day])
        actual = execute_causal_day(plan["buy"], loads[day], pvs[day], e_start)
        result = {
            **plan,
            "planned_charge": plan["charge"],
            "planned_discharge": plan["discharge"],
            "planned_energy": plan["energy"],
            "charge": actual["charge"],
            "discharge": actual["discharge"],
            "energy": actual["energy"],
            "emergency": actual["emergency"],
            "surplus": actual["surplus"],
            "fixed_emergency": fixed_emergency,
            "fixed_surplus": fixed_surplus,
        }
        checks = validate_day(
            result,
            loads[day],
            pvs[day],
            result["emergency"],
            result["surplus"],
        )
        max_checks["balance"] = max(max_checks["balance"], checks["balance"])
        max_checks["storage"] = max(max_checks["storage"], checks["storage"])
        max_checks["min_energy"] = min(max_checks["min_energy"], checks["min_energy"])
        max_checks["max_energy"] = max(max_checks["max_energy"], checks["max_energy"])

        result.update(
            {
                "forecast_load": load0,
                "forecast_pv": pv0,
                "alpha_load": alpha_l0,
                "alpha_pv": alpha_g0,
            }
        )
        all_results.append(result)
        residuals.append((loads[day] - pvs[day]) - net0)
        e_start = float(result["energy"][-1])

        if day >= REPORT_START:
            point_loads.append(load0)
            point_pvs.append(pv0)
            persistence_loads.append(loads[day - 1])
            persistence_pvs.append(pvs[day - 1])

    report = all_results[REPORT_START:]
    calibration = all_results[:REPORT_START]
    report_dates = dates[REPORT_START:]
    actual_load = loads[REPORT_START:]
    actual_pv = pvs[REPORT_START:]
    point_load = np.asarray(point_loads)
    point_pv = np.asarray(point_pvs)
    pers_load = np.asarray(persistence_loads)
    pers_pv = np.asarray(persistence_pvs)
    plan_kwh = sum(float(np.sum(r["buy"])) for r in report)
    emergency_kwh = sum(float(np.sum(r["emergency"])) for r in report)
    surplus_kwh = sum(float(np.sum(r["surplus"])) for r in report)
    plan_cost = sum(float(np.dot(prices, r["buy"])) for r in report)
    emergency_cost = sum(float(np.dot(5.0 * prices, r["emergency"])) for r in report)
    emergency_days = sum(float(np.sum(r["emergency"])) > 0.01 for r in report)
    simultaneous = sum(
        int(np.sum((r["charge"] > TOL) & (r["discharge"] > TOL))) for r in report
    )
    calibration_plan_cost = sum(float(np.dot(prices, r["buy"])) for r in calibration)
    calibration_emergency_cost = sum(
        float(np.dot(5.0 * prices, r["emergency"])) for r in calibration
    )
    fixed_emergency_kwh = sum(float(np.sum(r["fixed_emergency"])) for r in report)
    fixed_emergency_cost = sum(
        float(np.dot(5.0 * prices, r["fixed_emergency"])) for r in report
    )
    write_output(report_dates, prices, report)

    lines = [
        "Q2 FINAL MODEL METRICS",
        f"january_calibration_cost={calibration_plan_cost + calibration_emergency_cost:.4f} yuan",
        f"january_plan_cost={calibration_plan_cost:.4f} yuan",
        f"january_emergency_cost={calibration_emergency_cost:.4f} yuan",
        f"report_days={len(report)}",
        metric_line("adaptive_load", actual_load, point_load),
        metric_line("persistence_load", actual_load, pers_load),
        metric_line("adaptive_pv", actual_pv, point_pv),
        metric_line("persistence_pv", actual_pv, pers_pv),
        f"planned_purchase={plan_kwh:.4f} kWh",
        f"emergency_purchase={emergency_kwh:.4f} kWh",
        f"surplus_energy={surplus_kwh:.4f} kWh",
        f"planned_cost={plan_cost:.4f} yuan",
        f"emergency_cost={emergency_cost:.4f} yuan",
        f"total_cost={plan_cost + emergency_cost:.4f} yuan",
        f"fixed_schedule_emergency={fixed_emergency_kwh:.4f} kWh",
        f"fixed_schedule_emergency_cost={fixed_emergency_cost:.4f} yuan",
        f"fixed_schedule_total_cost={plan_cost + fixed_emergency_cost:.4f} yuan",
        f"emergency_days={emergency_days}/{len(report)}",
        f"ending_energy={report[-1]['energy'][-1]:.4f} kWh",
        f"simultaneous_charge_discharge_slots={simultaneous}",
        f"max_balance_residual={max_checks['balance']:.8g}",
        f"max_storage_residual={max_checks['storage']:.8g}",
        f"global_min_energy={max_checks['min_energy']:.4f} kWh",
        f"global_max_energy={max_checks['max_energy']:.4f} kWh",
        f"output={OUTPUT}",
    ]
    SUMMARY.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()

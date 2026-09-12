from pathlib import Path

import openpyxl

import q1_final


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "result1.xlsx"
OUTPUT = ROOT / "result1_补充修正版.xlsx"


def main():
    records, _ = q1_final.load_data()
    result = q1_final.build_and_solve(records)
    q1_final.validate(records, result)

    workbook = openpyxl.load_workbook(SOURCE)
    plan_sheet = workbook["计划购电量"]
    for row, value in enumerate(q1_final.source_order(result["buy"]), start=2):
        plan_sheet.cell(row=row, column=2, value=round(value, 4))

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        start = 24 * block
        stop = start + 24
        storage_sheet.cell(
            row=block + 2,
            column=2,
            value=round(sum(result["charge"][start:stop]), 4),
        )
        storage_sheet.cell(
            row=block + 2,
            column=3,
            value=round(sum(result["discharge"][start:stop]), 4),
        )
    storage_sheet.cell(row=2, column=5, value=round(result["energy"][0], 4))
    storage_sheet.cell(row=3, column=5, value=round(result["energy"][-1], 4))
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

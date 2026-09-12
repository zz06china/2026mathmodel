from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

import q3_final
import q3_sensitivity


HERE = Path(__file__).resolve().parent
OUTPUT_XLSX = HERE / "q3_selected_sensitivity.xlsx"
OUTPUT_TXT = HERE / "q3_selected_sensitivity.txt"


def main():
    prices, prior_load, _, dates, loads, pvs = q3_final.q2_final.load_inputs()
    cases = (
        ("窗口14天", 14, "band", 600.0, None),
        ("窗口56天", 56, "band", 600.0, None),
        ("终端区间5400-6600 kWh", 28, "band", 600.0, None),
        ("终端偏离软惩罚", 28, "soft", 600.0, None),
    )
    results = []
    for label, window, mode, band, penalty in cases:
        print(f"running {label}", flush=True)
        metrics = q3_sensitivity.run_case(
            prices,
            prior_load,
            dates,
            loads,
            pvs,
            q3_final.RELEASE_QUANTILES,
            window,
            mode,
            band,
            penalty,
        )
        results.append((label, window, mode, metrics))
        print(
            f"done total={metrics['total_cost']:.2f}, "
            f"emergency={metrics['emergency_kwh']:.2f} kWh, "
            f"time={metrics['elapsed_seconds']:.1f}s",
            flush=True,
        )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "最终模型敏感性补充"
    headers = [
        "参数方案",
        "调整后购电费/元",
        "紧急购电费/元",
        "总费用/元",
        "紧急购电量/kWh",
        "紧急购电天数",
        "期末SOC/kWh",
    ]
    sheet.append(headers)
    for label, _, _, metrics in results:
        sheet.append(
            [
                label,
                metrics["scheduled_cost"],
                metrics["emergency_cost"],
                metrics["total_cost"],
                metrics["emergency_kwh"],
                metrics["emergency_days"],
                metrics["end_energy"],
            ]
        )
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    for column, width in zip("ABCDEFG", (27, 19, 18, 18, 19, 16, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(OUTPUT_XLSX)

    lines = ["Q3 SELECTED MODEL SENSITIVITY", "q_release=0.80,0.70,0.70,0.70"]
    for label, window, mode, metrics in results:
        lines.append(
            f"case={label},window={window},terminal={mode},"
            f"scheduled_cost={metrics['scheduled_cost']:.6f},"
            f"emergency_cost={metrics['emergency_cost']:.6f},"
            f"total_cost={metrics['total_cost']:.6f},"
            f"emergency_kwh={metrics['emergency_kwh']:.6f},"
            f"emergency_days={metrics['emergency_days']},"
            f"end_energy={metrics['end_energy']:.6f}"
        )
    OUTPUT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"xlsx={OUTPUT_XLSX}")
    print(f"txt={OUTPUT_TXT}")


if __name__ == "__main__":
    main()

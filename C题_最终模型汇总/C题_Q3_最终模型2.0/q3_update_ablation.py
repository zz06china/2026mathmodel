from __future__ import annotations

from itertools import combinations
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

import q3_final


HERE = Path(__file__).resolve().parent
OUTPUT_TXT = HERE / "q3_update_ablation.txt"
OUTPUT_XLSX = HERE / "q3_update_ablation_results.xlsx"


def main():
    prices, prior_load, _, dates, loads, pvs = q3_final.q2_final.load_inputs()
    pv_forecasts, _ = q3_final.load_pv_forecasts(dates, pvs)
    net_load_margins = q3_final.build_net_load_margins(
        loads,
        pvs,
        pv_forecasts,
        prior_load,
    )

    choices = (36, 72, 108)
    cases = [combo for count in range(4) for combo in combinations(choices, count)]
    rows = []
    for combo in cases:
        releases = (0, *combo)
        label = "0" + "".join(f"+{slot // 6}" for slot in combo)
        print(f"running releases={label}", flush=True)
        results, checks = q3_final.simulate_strategy(
            f"releases={label}",
            prices,
            prior_load,
            dates,
            loads,
            pvs,
            pv_forecasts,
            net_load_margins,
            releases,
            True,
            terminal_mode=q3_final.FINAL_TERMINAL_MODE,
            terminal_band=q3_final.FINAL_TERMINAL_BAND,
            verbose=False,
        )
        values = q3_final.summarize(results)
        rows.append(
            {
                "label": label,
                "release_slots": releases,
                "scheduled_cost": values["scheduled_cost"],
                "emergency_cost": values["emergency_cost"],
                "total_cost": values["total_cost"],
                "emergency_kwh": values["emergency_kwh"],
                "emergency_days": values["emergency_days"],
                "daily_emergency_p95_kwh": values["daily_emergency_p95_kwh"],
                "emergency_cost_cvar95": values["daily_emergency_cost_cvar95"],
                "max_balance_residual": checks["balance"],
                "max_storage_residual": checks["storage"],
            }
        )
        print(
            f"done total={values['total_cost']:.2f}, "
            f"emergency={values['emergency_kwh']:.2f} kWh",
            flush=True,
        )

    rows.sort(key=lambda row: row["total_cost"])
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "预报时点消融"
    headers = [
        "调整预报时点",
        "调整后购电费/元",
        "紧急购电费/元",
        "总费用/元",
        "紧急购电量/kWh",
        "紧急购电天数",
        "日紧急购电量P95/kWh",
        "日紧急购电费CVaR95/元",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(
            [
                row["label"],
                row["scheduled_cost"],
                row["emergency_cost"],
                row["total_cost"],
                row["emergency_kwh"],
                row["emergency_days"],
                row["daily_emergency_p95_kwh"],
                row["emergency_cost_cvar95"],
            ]
        )
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    for column in range(1, sheet.max_column + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 22
    sheet.freeze_panes = "A2"
    lines = ["Q3 UPDATE-TIME ABLATION", "sorted_by=total_cost"]
    for row in rows:
        lines.append(
            f"releases={row['label']},"
            f"scheduled_cost={row['scheduled_cost']:.6f},"
            f"emergency_cost={row['emergency_cost']:.6f},"
            f"total_cost={row['total_cost']:.6f},"
            f"emergency_kwh={row['emergency_kwh']:.6f},"
            f"emergency_days={row['emergency_days']},"
            f"daily_emergency_p95_kwh={row['daily_emergency_p95_kwh']:.6f},"
            f"daily_emergency_cost_cvar95={row['emergency_cost_cvar95']:.6f}"
        )
    OUTPUT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    workbook.save(OUTPUT_XLSX)
    print(f"best_releases={rows[0]['label']}")
    print(f"xlsx={OUTPUT_XLSX}")
    print(f"txt={OUTPUT_TXT}")


if __name__ == "__main__":
    main()

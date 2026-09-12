from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
RESULT = HERE / "result3.xlsx"
METRICS = HERE / "q3_metrics.txt"
OUTPUT = HERE / "第三问最终模型.md"
TARGET_DATES = [
    datetime(2025, 3, 20),
    datetime(2025, 6, 21),
    datetime(2025, 9, 23),
    datetime(2025, 12, 21),
]
PLAN_TIMES = [
    "10:00-10:10",
    "12:00-12:10",
    "14:00-14:10",
    "16:00-16:10",
    "18:00-18:10",
    "20:00-20:10",
]


def read_metrics():
    sections = {}
    current = None
    for raw in METRICS.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("[") and line.endswith("]"):
            current = line[1:-1]
            sections[current] = {}
        elif current and "=" in line:
            key, value = line.split("=", 1)
            try:
                sections[current][key] = float(value)
            except ValueError:
                sections[current][key] = value
    return sections


def read_tables():
    wb = openpyxl.load_workbook(RESULT, data_only=True, read_only=False)
    plan_tables = {}
    for sheet_name in ("计划购电量", "调整购电量"):
        sheet = wb[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        rows = []
        for target in TARGET_DATES:
            row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(r, 1).value == target)
            values = [float(sheet.cell(row, headers.index(label) + 1).value) for label in PLAN_TIMES]
            rows.append(
                [target.strftime("%Y-%m-%d"), *values, float(sheet.cell(row, 146).value), float(sheet.cell(row, 147).value)]
            )
        plan_tables[sheet_name] = rows

    storage_sheet = wb["充放电量"]
    storage_rows = []
    for target in TARGET_DATES:
        row = next(r for r in range(2, storage_sheet.max_row + 1) if storage_sheet.cell(r, 1).value == target)
        for block in range(6):
            storage_rows.append(
                [
                    target.strftime("%Y-%m-%d") if block == 0 else "",
                    storage_sheet.cell(row + block, 2).value,
                    float(storage_sheet.cell(row + block, 3).value),
                    float(storage_sheet.cell(row + block, 4).value),
                    float(storage_sheet.cell(row, 6).value) if block == 0 else (
                        float(storage_sheet.cell(row + 1, 6).value) if block == 1 else None
                    ),
                ]
            )

    emergency_sheet = wb["紧急购电量"]
    emergency_rows = []
    for target in TARGET_DATES:
        first = next(
            (r for r in range(2, emergency_sheet.max_row + 1) if emergency_sheet.cell(r, 1).value == target),
            None,
        )
        if first is None:
            emergency_rows.append([target.strftime("%Y-%m-%d"), "无", 0.0])
            continue
        stop = next(
            (r for r in range(first + 1, emergency_sheet.max_row + 1) if emergency_sheet.cell(r, 1).value is not None),
            emergency_sheet.max_row + 1,
        )
        for row in range(first, stop):
            emergency_rows.append(
                [
                    target.strftime("%Y-%m-%d") if row == first else "",
                    emergency_sheet.cell(row, 2).value,
                    float(emergency_sheet.cell(row, 3).value),
                ]
            )
    wb.close()
    return plan_tables, storage_rows, emergency_rows


def md_table(headers, rows, formats=None):
    formats = formats or [str] * len(headers)
    lines = ["| " + " | ".join(headers) + " |", "|" + "|".join(["---"] * len(headers)) + "|"]
    for row in rows:
        values = []
        for value, formatter in zip(row, formats):
            if value is None:
                values.append("")
            elif callable(formatter):
                values.append(formatter(value))
            else:
                values.append(str(value))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def build():
    metrics = read_metrics()
    plan_tables, storage_rows, emergency_rows = read_tables()
    comparison_order = ["static_robust", "robust_06", "robust_0612", "nominal_mpc", "robust_mpc"]
    comparison_labels = {
        "static_robust": "仅0:00预报（鲁棒）",
        "robust_06": "增加6:00更新",
        "robust_0612": "增加6:00、12:00更新",
        "nominal_mpc": "四时刻普通MPC",
        "robust_mpc": "四时刻净负荷分位数鲁棒MPC",
    }
    comparison_rows = []
    for name in comparison_order:
        item = metrics[name]
        comparison_rows.append(
            [
                comparison_labels[name],
                item["scheduled_cost"] / 1e4,
                item["emergency_cost"] / 1e4,
                item["total_cost"] / 1e4,
                item["emergency_kwh"],
                int(item["emergency_days"]),
            ]
        )
    ablation_rows = [
        ["0:00、6:00、12:00、18:00", 1369.8765, 58551.8747],
        ["0:00、12:00、18:00", 1376.7254, 66174.3700],
        ["0:00、6:00、12:00", 1389.2040, 97948.2944],
        ["0:00、6:00、18:00", 1389.8114, 58523.1800],
        ["0:00、12:00", 1395.4343, 104199.9100],
        ["0:00、18:00", 1400.7637, 66008.1500],
        ["仅0:00", 1418.2946, 101214.7395],
        ["0:00、6:00", 1423.8972, 129087.7652],
    ]

    selected = metrics["robust_mpc"]
    info = metrics["information_value"]
    headers = ["日期", "10:00", "12:00", "14:00", "16:00", "18:00", "20:00", "全天/kWh", "费用/元"]
    table_formats = [str] + [lambda x: f"{x:.4f}"] * 8
    storage_formats = [str, str, lambda x: f"{x:.4f}", lambda x: f"{x:.4f}", lambda x: f"{x:.4f}"]
    emergency_formats = [str, str, lambda x: f"{x:.4f}"]
    fmt4 = lambda x: f"{x:.4f}"

    text = fr"""# 问题三最终模型

## 1. 模型决策

第三问最终采用 **非对称结算约束下的分阶段净负荷分位数鲁棒模型预测控制（Stagewise Net-load Quantile-Robust MPC）**。模型同时保留三项题意核心：每天 0:00 形成基准计划购电量，6:00、12:00 和 18:00 只能调整尚未执行的购电量，实际运行时剩余缺口按同时段 5 倍电价紧急购入。

不采用“逐时段取绝对最坏值”的区间鲁棒模型，也不再只修正光伏预测。本文对 **完整净负荷预测残差** 计算同发布时刻、同预测步长的单侧经验分位数，使裕量同时覆盖负荷偏高与光伏偏低两类风险；0:00 与调整阶段分别采用 0.80 和 0.70 分位数，在经济性与可靠性之间形成可解释的折中。

## 2. 问题分析

问题二只有每天 0:00 一次计划决策，问题三则增加三次光伏预报更新和非对称调整结算。因而第三问不能仅把光伏预测值替换进问题二模型，而应区分：

1. **基准计划层**：0:00 根据当时可得信息制定全天计划购电量；
2. **滚动调整层**：6:00、12:00、18:00 利用新预报和当前实际 SOC 重算未来 24 h，但只执行至下一发布时刻；
3. **因果执行层**：每个 10 min 时段到达后，根据当期实际负荷、实际光伏和当前 SOC 执行固定反馈规则，无法覆盖的缺口才形成紧急购电。

这种分层既允许题目规定的策略调整，又不会把未来真实数据用于当前决策。

## 3. 模型假设

1. 电价按附件 1 给定的日内曲线逐日重复，每个 10 min 时段内保持不变。
2. 附件 3 的 24 个整点功率预报在相邻整点间作线性插值，并乘以 $\Delta t=1/6$ h 转换为时段电量。
3. 负荷预测沿用问题二的滞后 1 日与滞后 7 日自适应融合方法；权重仅由发布时刻以前的数据确定。
4. 0:00 基准计划一经形成即作为调整结算基准；后续只能修改尚未执行的时段。
5. 每个时段最终采用最近一次发布时刻生成的购电量，历史已执行决策不可修改。
6. 储能是用户侧内部调节设备，可依据当前观测实时反馈，但反馈规则在运行前已确定，且不调用未来真实值。
7. 暂不计储能衰减、启停费用和固定运维费用；弃光不产生收益。

## 4. 符号说明

| 符号 | 含义 | 单位 |
|---|---|---|
| $d,t$ | 日期及 10 min 时段编号 | — |
| $r\in\{{0,6,12,18\}}$ | 光伏预报发布时刻 | h |
| $L_{{d,t}},G_{{d,t}}$ | 实际负荷和实际光伏电量 | kWh |
| $\widehat L^r_{{d,t}},\widehat G^r_{{d,t}}$ | 发布时刻 $r$ 可得的负荷、光伏预测 | kWh |
| $e^r_{{d,t}},\Delta^r_{{d,t}}$ | 净负荷预测残差及单侧风险裕量 | kWh |
| $B^0_{{d,t}}$ | 0:00 制定的基准计划购电量 | kWh |
| $B^r_{{d,t}}$ | 发布时刻 $r$ 对未来时段制定的购电量 | kWh |
| $\bar B_{{d,t}}$ | 最终执行的计划或调整购电量 | kWh |
| $U_{{d,t}},V_{{d,t}}$ | 相对基准计划的增加量和减少量 | kWh |
| $C_{{d,t}},D_{{d,t}}$ | 储能交流侧充电量和放电量 | kWh |
| $E_{{d,t}}$ | 时段起点储电量 | kWh |
| $R_{{d,t}},S_{{d,t}}$ | 紧急购电量和富余电量 | kWh |
| $p_t$ | 正常交易电价 | 元/kWh |

## 5. 数据预处理与风险裕量

### 5.1 时间和量纲统一

附件 2 的 10 min 功率按

$$L_{{d,t}}=P^L_{{d,t}}\Delta t,\qquad G_{{d,t}}=P^G_{{d,t}}\Delta t,\qquad \Delta t=\frac16\text{{ h}}$$

转换为电量。附件 3 每次发布的第 $h$ 小时整点预测，利用发布前最后一个已观测光伏功率作为零时刻锚点，对 10 min 时段中心作线性插值。该锚点不使用发布之后的数据。

### 5.2 负荷预测

对目标日负荷采用

$$\widehat L_{{d,t}}=\alpha_dL_{{d-1,t}}+(1-\alpha_d)L_{{d-7,t}},$$

其中 $\alpha_d\in\{{0,0.1,\ldots,1\}}$，依据计划日前最多 56 天历史样本的绝对误差滚动选择。

### 5.3 分阶段净负荷分位数裕量

定义发布时刻 $r$ 对预测步长 $k$ 的完整净负荷残差

$$e^r_{{j,k}}=(L_{{j,k}}-G_{{j,k}})-(\widehat L^r_{{j,k}}-\widehat G^r_{{j,k}}).$$

根据基准购电和上调购电的不同边际价格，设置

$$q_0=0.80,\qquad q_6=q_{{12}}=q_{{18}}=0.70.$$

针对每个发布时刻和预测步长，仅用此前 28 天已经实现的同类残差计算

$$\Delta^r_{{d,k}}=\max\left\{{0,Q_{{q_r}}\left(e^r_{{d-28:d-1,k}}\right)\right\}},$$

从而构造鲁棒净负荷

$$N^{{rob,r}}_{{d,k}}=\widehat L^r_{{d,k}}-\widehat G^r_{{d,k}}+\Delta^r_{{d,k}}.$$

0:00 正常购电的临界分位数为 $1-p/(5p)=0.8$；调整阶段上调购电按 $1.5p$ 结算，对应 $1-1.5p/(5p)=0.7$。该参数由题目价格比确定，历史窗口与当日运行严格分离，不利用未来残差事后调参。

## 6. 分阶段净负荷分位数鲁棒 MPC 模型

### 6.1 0:00 基准计划

以当前实际储电量为初值，在未来 24 h 预测域内求解

$$\min Z^0_d=\sum_k p_k B^0_{{d,k}}+\varepsilon\sum_k(C_{{d,k}}+D_{{d,k}}+S_{{d,k}}),$$

其中 $\varepsilon=10^{{-8}}$ 仅用于从等价最优解中筛除无意义能量循环，不改变费用精度。

### 6.2 发布时刻的调整结算

对尚未执行的时段设置

$$B^r_{{d,t}}-B^0_{{d,t}}=U^r_{{d,t}}-V^r_{{d,t}},\qquad U^r_{{d,t}},V^r_{{d,t}}\ge0.$$

若调整量高于计划量，增加部分按 $1.5p_t$ 结算；若低于计划量，取消部分仍承担 $0.5p_t$ 的违约成本。因此单时段计划与调整结算费用可写为

$$c^{{sch}}_{{d,t}}=p_tB^0_{{d,t}}+1.5p_tU^r_{{d,t}}-0.5p_tV^r_{{d,t}}.$$

当 $B^r<B^0$ 时，上式等价于 $p_tB^r+0.5p_t(B^0-B^r)$；当 $B^r>B^0$ 时，等价于 $p_tB^0+1.5p_t(B^r-B^0)$，与题目两种调整规则一致。

### 6.3 预测域约束

每次优化满足

$$B^r_{{d,k}}+D^r_{{d,k}}=N^{{rob,r}}_{{d,k}}+C^r_{{d,k}}+S^r_{{d,k}},$$

$$E^r_{{d,k+1}}=E^r_{{d,k}}+0.9C^r_{{d,k}}-\frac{{D^r_{{d,k}}}}{{0.9}},$$

$$1200\le E^r_{{d,k}}\le10800,$$

$$0\le C^r_{{d,k}},D^r_{{d,k}}\le833.3333,\qquad B^r_{{d,k}},S^r_{{d,k}}\ge0.$$

预测域末端采用带状约束

$$5400\le E^r_{{d,N}}\le6600,$$

既防止滚动优化在视野末端透支储能，又允许模型利用预测域末端附近的有效峰谷信息，避免硬性回到 6000 kWh 所造成的边界调度。

### 6.4 因果执行与紧急购电

每个 10 min 时段到达后，令

$$n_{{d,t}}=L_{{d,t}}-G_{{d,t}}-\bar B_{{d,t}}.$$

若 $n_{{d,t}}\ge0$，则

$$D_{{d,t}}=\min\left\{{n_{{d,t}},833.3333,0.9(E_{{d,t}}-1200)\right\}},$$

$$R_{{d,t}}=n_{{d,t}}-D_{{d,t}}.$$

若 $n_{{d,t}}<0$，则

$$C_{{d,t}}=\min\left\{{-n_{{d,t}},833.3333,\frac{{10800-E_{{d,t}}}}{{0.9}}\right\}},$$

$$S_{{d,t}}=-n_{{d,t}}-C_{{d,t}}.$$

由分段规则可知同一时段不会同时充放电。实际期末 SOC 结转到下一发布时刻和下一日。

### 6.5 总费用

最终总费用为

$$Z=\sum_{{d,t}}\left[p_tB^0_{{d,t}}+1.5p_tU_{{d,t}}-0.5p_tV_{{d,t}}+5p_tR_{{d,t}}\right].$$

## 7. 滚动求解流程

1. 0:00 使用 0:00 预报和当前 SOC 求解 24 h 基准计划，并执行 0:00—6:00；
2. 6:00 读取新预报与实际 SOC，仅重算未执行时段，执行 6:00—12:00；
3. 12:00 同理执行至 18:00；
4. 18:00 同理执行至 24:00；
5. 每段内部按因果反馈规则运行，记录最终调整购电、储能和紧急购电结果；
6. 将 24:00 实际 SOC 结转到下一天。

## 8. 全年结果与模型比较

{md_table(["策略", "调整后购电费/万元", "紧急购电费/万元", "总费用/万元", "紧急购电量/kWh", "紧急购电天数"], comparison_rows, [str, fmt4, fmt4, fmt4, fmt4, str])}

最终采用的四时刻净负荷分位数鲁棒 MPC 在 2025 年 2 月 1 日至 12 月 31 日的调整后购电费为 **{selected['scheduled_cost']:.2f} 元**，紧急购电费为 **{selected['emergency_cost']:.2f} 元**，总费用为 **{selected['total_cost']:.2f} 元**，紧急购电量为 **{selected['emergency_kwh']:.4f} kWh**。

相较仅使用 0:00 预报的同风险口径模型，总费用减少 **{info['total_cost_saving']:.2f} 元**，紧急购电量减少 **{info['emergency_kwh_reduction']:.4f} kWh**。相较四时刻普通 MPC，总费用减少 **{metrics['nominal_mpc']['total_cost']-selected['total_cost']:.2f} 元**，紧急购电量减少 **{metrics['nominal_mpc']['emergency_kwh']-selected['emergency_kwh']:.4f} kWh**，说明风险裕量并非仅增加保守购电，而是降低了高价紧急购电后的实际总成本。

### 8.1 更新时点全组合消融

{md_table(["使用的预报时刻", "总费用/万元", "紧急购电量/kWh"], ablation_rows, [str, fmt4, fmt4])}

四时刻方案取得最低总费用，其紧急购电量仅比全组合中的最低值多 28.69 kWh。6:00 单独加入时并不占优，但与 12:00、18:00 联用后会通过 SOC 状态和后续调整结算产生协同作用，因此最终保留题目提供的全部四次预报。

### 8.2 尾部风险检验

最终方案日紧急购电量的 95% 分位数为 **{selected['daily_emergency_p95_kwh']:.2f} kWh**，单时段最大紧急购电量为 **{selected['max_interval_emergency_kwh']:.2f} kWh**，全年紧急购电持续时间为 **{selected['emergency_duration_hours']:.2f} h**，日紧急购电费 CVaR$_{{0.95}}$ 为 **{selected['daily_emergency_cost_cvar95']:.2f} 元**。这些指标补充了均值费用，说明低总成本并非依靠把风险集中到少数极端日期获得。

## 9. 指定日期结果

### 9.1 0:00 计划购电量

{md_table(headers, plan_tables['计划购电量'], table_formats)}

### 9.2 最终调整购电量

{md_table(headers, plan_tables['调整购电量'], table_formats)}

### 9.3 储能运行结果

{md_table(["日期", "时间段", "充电量/kWh", "放电量/kWh", "SOC（0:00/24:00）"], storage_rows, storage_formats)}

### 9.4 紧急购电结果

{md_table(["日期", "紧急购电时间段", "紧急购电量/kWh"], emergency_rows, emergency_formats)}

## 10. 可行性与模型评价

最终策略的最大供需平衡残差为 **{selected['max_balance_residual']:.3e} kWh**，最大储能递推残差为 **{selected['max_storage_residual']:.3e} kWh**；全年 SOC 均位于 1200—10800 kWh，未出现同时充放电。

模型的优点是：准确体现非对称调整价格；严格保持信息因果性；以完整净负荷残差覆盖两类预测风险；利用四次预报、实际 SOC 和终端带状约束提升滚动决策质量；仍保持为可稳定全年求解的线性规划。其局限是逐预测步长分位数尚未显式刻画连续阴雨误差的轨迹相关性，也未计入电池寿命成本；若需更强极端风险保障，可进一步引入误差轨迹场景与 CVaR。
"""
    OUTPUT.write_text(text, encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    build()

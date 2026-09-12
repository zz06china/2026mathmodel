from __future__ import annotations

from pathlib import Path
from time import perf_counter

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

import q3_final


HERE = Path(__file__).resolve().parent
OUTPUT_XLSX = HERE / "q3_sensitivity.xlsx"
OUTPUT_TXT = HERE / "q3_sensitivity.txt"

QUANTILES = (0.70, 0.75, 0.80, 0.85, 0.90)
WINDOWS = (14, 28, 56)
TERMINAL_CASES = (
    ("固定6000 kWh", "fixed", 600.0, None),
    ("区间5400-6600 kWh", "band", 600.0, None),
    ("偏离6000 kWh软惩罚", "soft", 600.0, None),
)


def run_case(
    prices,
    prior_load,
    dates,
    loads,
    pvs,
    quantile,
    window,
    terminal_mode=q3_final.FINAL_TERMINAL_MODE,
    terminal_band=600.0,
    terminal_penalty=None,
):
    if np.isscalar(quantile):
        quantile_label = f"{float(quantile):.2f}"
    else:
        quantile_label = "-".join(f"{float(value):.2f}" for value in quantile)
    pv_forecasts, _ = q3_final.load_pv_forecasts(
        dates,
        pvs,
    )
    net_load_margins = q3_final.build_net_load_margins(
        loads,
        pvs,
        pv_forecasts,
        prior_load,
        error_window=window,
        risk_quantile=quantile,
    )
    started = perf_counter()
    results, checks = q3_final.simulate_strategy(
        f"q={quantile_label},w={window},{terminal_mode}",
        prices,
        prior_load,
        dates,
        loads,
        pvs,
        pv_forecasts,
        net_load_margins,
        q3_final.RELEASE_SLOTS,
        True,
        terminal_mode=terminal_mode,
        terminal_band=terminal_band,
        terminal_penalty=terminal_penalty,
        verbose=False,
    )
    summary = q3_final.summarize(results)
    summary["max_balance_residual"] = checks["balance"]
    summary["max_storage_residual"] = checks["storage"]
    summary["elapsed_seconds"] = perf_counter() - started
    return summary


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    widths = [22, 15, 18, 18, 18, 16, 16]
    for index, width in enumerate(widths[: len(headers)], start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def row_from(label, metrics):
    return [
        label,
        metrics["scheduled_cost"],
        metrics["emergency_cost"],
        metrics["total_cost"],
        metrics["emergency_kwh"],
        metrics["emergency_days"],
        metrics["end_energy"],
    ]


def main():
    prices, prior_load, _, dates, loads, pvs = q3_final.q2_final.load_inputs()
    cache = {}

    def get_case(
        q,
        window,
        mode=q3_final.FINAL_TERMINAL_MODE,
        band=q3_final.FINAL_TERMINAL_BAND,
        penalty=None,
    ):
        key = (q, window, mode, band, penalty)
        if key not in cache:
            print(f"running q={q:.2f}, window={window}, terminal={mode}", flush=True)
            cache[key] = run_case(
                prices,
                prior_load,
                dates,
                loads,
                pvs,
                q,
                window,
                mode,
                band,
                penalty,
            )
            print(
                f"done total={cache[key]['total_cost']:.2f}, "
                f"emergency={cache[key]['emergency_kwh']:.2f} kWh, "
                f"time={cache[key]['elapsed_seconds']:.1f}s",
                flush=True,
            )
        return cache[key]

    quantile_rows = [row_from(f"q={q:.2f}", get_case(q, 28)) for q in QUANTILES]
    window_rows = [row_from(f"窗口={w}天", get_case(0.80, w)) for w in WINDOWS]
    terminal_rows = []
    for label, mode, band, penalty in TERMINAL_CASES:
        terminal_rows.append(row_from(label, get_case(0.80, 28, mode, band, penalty)))

    headers = [
        "参数方案",
        "调整后购电费/元",
        "紧急购电费/元",
        "总费用/元",
        "紧急购电量/kWh",
        "紧急购电天数",
        "期末SOC/kWh",
    ]
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "分位数敏感性", headers, quantile_rows)
    add_sheet(workbook, "窗口敏感性", headers, window_rows)
    add_sheet(workbook, "终端SOC敏感性", headers, terminal_rows)
    workbook.save(OUTPUT_XLSX)

    baseline = get_case(0.80, 28)
    lines = [
        "Q3 SENSITIVITY ANALYSIS",
        "All cases use causal four-release robust MPC on 2025-02-01 to 2025-12-31.",
        "",
        "[quantile]",
    ]
    for q, row in zip(QUANTILES, quantile_rows):
        lines.append(
            f"q={q:.2f},scheduled_cost={row[1]:.6f},emergency_cost={row[2]:.6f},"
            f"total_cost={row[3]:.6f},emergency_kwh={row[4]:.6f},emergency_days={row[5]}"
        )
    lines.extend(["", "[window]"])
    for window, row in zip(WINDOWS, window_rows):
        lines.append(
            f"window={window},scheduled_cost={row[1]:.6f},emergency_cost={row[2]:.6f},"
            f"total_cost={row[3]:.6f},emergency_kwh={row[4]:.6f},emergency_days={row[5]}"
        )
    lines.extend(["", "[terminal]"])
    for case, row in zip(TERMINAL_CASES, terminal_rows):
        lines.append(
            f"policy={case[0]},scheduled_cost={row[1]:.6f},emergency_cost={row[2]:.6f},"
            f"total_cost={row[3]:.6f},emergency_kwh={row[4]:.6f},"
            f"emergency_days={row[5]},end_energy={row[6]:.6f}"
        )
    lines.extend(
        [
            "",
            "[baseline]",
            "quantile=0.80",
            "window=28",
            "terminal=band_5400_6600",
            f"total_cost={baseline['total_cost']:.6f}",
            f"max_balance_residual={baseline['max_balance_residual']:.12g}",
            f"max_storage_residual={baseline['max_storage_residual']:.12g}",
        ]
    )
    OUTPUT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"xlsx={OUTPUT_XLSX}")
    print(f"txt={OUTPUT_TXT}")


if __name__ == "__main__":
    main()

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "第三问优化与敏感性分析汇总.xlsx"


def parse_records(path):
    records = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if "=" not in line or "," not in line:
            continue
        record = {}
        for item in line.split(","):
            if "=" not in item:
                continue
            key, value = item.split("=", 1)
            record[key] = value
        if "total_cost" in record:
            records.append(record)
    return records


def number(record, key):
    return float(record[key])


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = 20
    sheet.column_dimensions["A"].width = 27
    sheet.freeze_panes = "A2"


def main():
    import build_q3_model_report as report_builder

    metrics = report_builder.read_metrics()
    strategy_names = [
        ("static_robust", "仅0:00净负荷鲁棒"),
        ("robust_06", "0:00、6:00"),
        ("robust_0612", "0:00、6:00、12:00"),
        ("nominal_mpc", "四时刻普通MPC"),
        ("robust_mpc", "四时刻净负荷分位数鲁棒MPC（采用）"),
    ]
    cost_headers = [
        "方案",
        "调整后购电费/万元",
        "紧急购电费/万元",
        "总费用/万元",
        "紧急购电量/kWh",
        "紧急购电天数",
    ]
    strategy_rows = []
    for key, label in strategy_names:
        item = metrics[key]
        strategy_rows.append([
            label,
            item["scheduled_cost"] / 10000,
            item["emergency_cost"] / 10000,
            item["total_cost"] / 10000,
            item["emergency_kwh"],
            int(item["emergency_days"]),
        ])

    ablation_rows = []
    for record in parse_records(HERE / "q3_update_ablation.txt"):
        ablation_rows.append([
            record["releases"],
            number(record, "total_cost") / 10000,
            number(record, "emergency_kwh"),
        ])

    selected = metrics["robust_mpc"]
    risk_rows = [[
        "四时刻净负荷分位数鲁棒MPC",
        selected["daily_emergency_p95_kwh"],
        selected["max_interval_emergency_kwh"],
        selected["emergency_duration_hours"],
        selected["daily_emergency_cost_var95"],
        selected["daily_emergency_cost_cvar95"],
    ]]
    parameter_rows = [
        ["发布时刻", "0:00、6:00、12:00、18:00"],
        ["风险变量", "完整净负荷预测残差"],
        ["历史窗口", "28天"],
        ["分位数", "0.80、0.70、0.70、0.70"],
        ["预测域", "固定24 h"],
        ["终端SOC", "5400—6600 kWh"],
    ]

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "策略比较", cost_headers, strategy_rows)
    add_sheet(
        workbook,
        "更新时点全组合",
        ["发布时刻", "总费用/万元", "紧急购电量/kWh"],
        ablation_rows,
    )
    add_sheet(
        workbook,
        "尾部风险",
        [
            "方案",
            "日紧急购电量P95/kWh",
            "单时段最大紧急购电量/kWh",
            "紧急购电持续时间/h",
            "日紧急购电费VaR95/元",
            "日紧急购电费CVaR95/元",
        ],
        risk_rows,
    )
    add_sheet(workbook, "最终参数", ["参数", "取值"], parameter_rows)
    workbook.save(OUTPUT)
    print(f"saved={OUTPUT}")


if __name__ == "__main__":
    main()

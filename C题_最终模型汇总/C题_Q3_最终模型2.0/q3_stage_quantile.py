from pathlib import Path
from time import perf_counter

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

import q3_final


HERE = Path(__file__).resolve().parent
OUTPUT_XLSX = HERE / "q3_stage_quantile.xlsx"
OUTPUT_TXT = HERE / "q3_stage_quantile.txt"
ADJUSTMENT_QUANTILES = (0.65, 0.70, 0.75)


def run_case(prices, prior_load, dates, loads, pvs, adjustment_quantile):
    release_quantiles = (0.80,) + (adjustment_quantile,) * 3
    pv_forecasts, _ = q3_final.load_pv_forecasts(
        dates,
        pvs,
    )
    net_load_margins = q3_final.build_net_load_margins(
        loads,
        pvs,
        pv_forecasts,
        prior_load,
        error_window=28,
        risk_quantile=release_quantiles,
    )
    started = perf_counter()
    results, checks = q3_final.simulate_strategy(
        f"staged_q_{adjustment_quantile:.2f}",
        prices,
        prior_load,
        dates,
        loads,
        pvs,
        pv_forecasts,
        net_load_margins,
        q3_final.RELEASE_SLOTS,
        True,
        terminal_mode=q3_final.FINAL_TERMINAL_MODE,
        terminal_band=q3_final.FINAL_TERMINAL_BAND,
        verbose=False,
    )
    metrics = q3_final.summarize(results)
    metrics["max_balance_residual"] = checks["balance"]
    metrics["max_storage_residual"] = checks["storage"]
    metrics["elapsed_seconds"] = perf_counter() - started
    return metrics


def main():
    prices, prior_load, _, dates, loads, pvs = q3_final.q2_final.load_inputs()
    rows = []
    details = []
    for q_adjust in ADJUSTMENT_QUANTILES:
        print(f"running q0=0.80, q_adjust={q_adjust:.2f}", flush=True)
        metrics = run_case(prices, prior_load, dates, loads, pvs, q_adjust)
        details.append(metrics)
        rows.append(
            [
                0.80,
                q_adjust,
                metrics["scheduled_cost"],
                metrics["emergency_cost"],
                metrics["total_cost"],
                metrics["emergency_kwh"],
                metrics["emergency_days"],
                metrics["end_energy"],
            ]
        )
        print(
            f"done total={metrics['total_cost']:.2f}, "
            f"emergency={metrics['emergency_kwh']:.2f} kWh, "
            f"time={metrics['elapsed_seconds']:.1f}s",
            flush=True,
        )

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "分阶段分位数"
    headers = [
        "0:00分位数",
        "调整阶段分位数",
        "调整后购电费/元",
        "紧急购电费/元",
        "总费用/元",
        "紧急购电量/kWh",
        "紧急购电天数",
        "期末SOC/kWh",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    for column, width in zip("ABCDEFGH", (14, 18, 19, 18, 18, 19, 16, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(OUTPUT_XLSX)

    lines = ["Q3 STAGE-SPECIFIC QUANTILE ANALYSIS"]
    for q_adjust, metrics in zip(ADJUSTMENT_QUANTILES, details):
        lines.append(
            f"q0=0.80,q_adjust={q_adjust:.2f},"
            f"scheduled_cost={metrics['scheduled_cost']:.6f},"
            f"emergency_cost={metrics['emergency_cost']:.6f},"
            f"total_cost={metrics['total_cost']:.6f},"
            f"emergency_kwh={metrics['emergency_kwh']:.6f},"
            f"emergency_days={metrics['emergency_days']},"
            f"end_energy={metrics['end_energy']:.6f},"
            f"max_balance_residual={metrics['max_balance_residual']:.12g},"
            f"max_storage_residual={metrics['max_storage_residual']:.12g}"
        )
    OUTPUT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"xlsx={OUTPUT_XLSX}")
    print(f"txt={OUTPUT_TXT}")


if __name__ == "__main__":
    main()

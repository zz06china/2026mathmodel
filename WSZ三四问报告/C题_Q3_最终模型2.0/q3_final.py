from __future__ import annotations

import argparse
from copy import copy
from datetime import datetime, time
from pathlib import Path
from shutil import copyfile
import sys

import numpy as np
import openpyxl
import pulp


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
Q2_DIR = ROOT / "C题_Q2_最终模型"
ATTACHMENT3 = ROOT / "C题" / "附件" / "附件3.xlsx"
TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result3.xlsx"
OUTPUT = HERE / "result3_optimized.xlsx"
METRICS = HERE / "q3_metrics.txt"

sys.path.insert(0, str(Q2_DIR))
import q2_final  # noqa: E402


DT = 1.0 / 6.0
N = 144
BLOCK = 36
RELEASE_SLOTS = (0, 36, 72, 108)
ETA_C = 0.9
ETA_D = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
Q_MAX = 5000.0 * DT
REPORT_START = 31
ERROR_WINDOW = 28
RISK_QUANTILE = 0.80
RELEASE_QUANTILES = (0.80, 0.70, 0.70, 0.70)
TOL = 1e-6
EXPORT_TOL = 5e-5
SOLVER = pulp.PULP_CBC_CMD(msg=False)


def load_pv_forecasts(
    dates,
    actual_pv,
    error_window=ERROR_WINDOW,
    risk_quantile=RELEASE_QUANTILES,
):
    if np.isscalar(risk_quantile):
        release_quantiles = np.full(len(RELEASE_SLOTS), float(risk_quantile))
    else:
        release_quantiles = np.asarray(risk_quantile, dtype=float)
        if release_quantiles.shape != (len(RELEASE_SLOTS),):
            raise ValueError("risk_quantile必须是标量或与四个发布时刻对应的长度4序列")
    if np.any((release_quantiles < 0.0) | (release_quantiles > 1.0)):
        raise ValueError("risk_quantile必须位于[0,1]")

    wb = openpyxl.load_workbook(ATTACHMENT3, data_only=True, read_only=True)
    sheet = wb.active
    raw = {}
    current_date = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        date_value, release_value = row[:2]
        if date_value not in (None, ""):
            current_date = datetime.strptime(str(date_value), "%Y-%m-%d").date()
        if current_date is None or release_value in (None, ""):
            continue
        release_hour = int(str(release_value).split(":")[0])
        values = np.asarray(row[2:26], dtype=float)
        if len(values) != 24:
            raise ValueError("附件3的单次预报长度不是24小时")
        raw[(current_date, release_hour)] = values
    wb.close()

    actual_power_flat = (actual_pv / DT).reshape(-1)
    forecasts = np.zeros((len(dates), 4, N))
    errors = np.full_like(forecasts, np.nan)
    query_hours = (np.arange(N) + 0.5) / 6.0
    interpolation_hours = np.arange(25, dtype=float)

    for day, date_value in enumerate(dates):
        date_key = date_value.date() if hasattr(date_value, "date") else date_value
        for release_index, release_slot in enumerate(RELEASE_SLOTS):
            release_hour = release_slot // 6
            hourly = raw[(date_key, release_hour)]
            issue_global = day * N + release_slot
            anchor_index = max(issue_global - 1, 0)
            anchor = actual_power_flat[anchor_index]
            power = np.interp(query_hours, interpolation_hours, np.r_[anchor, hourly])
            forecasts[day, release_index] = np.maximum(power, 0.0)
            available = min(N, len(actual_power_flat) - issue_global)
            if available > 0:
                errors[day, release_index, :available] = (
                    power[:available] - actual_power_flat[issue_global : issue_global + available]
                )

    margins = np.zeros_like(forecasts)
    for day in range(len(dates)):
        start = max(0, day - int(error_window))
        if day == start:
            continue
        history = errors[start:day]
        with np.errstate(all="ignore"):
            for release_index, quantile in enumerate(release_quantiles):
                margin = np.nanquantile(
                    history[:, release_index, :],
                    quantile,
                    axis=0,
                )
                margins[day, release_index] = np.maximum(
                    np.nan_to_num(margin, nan=0.0),
                    0.0,
                )
    return forecasts * DT, margins * DT


def load_forecast_horizon(loads, day, release_slot, prior_load):
    today, _ = q2_final.point_forecast(loads, day, 0, prior_load)
    tomorrow, _ = q2_final.point_forecast(loads, day, 1, prior_load)
    if release_slot == 0:
        return today.copy()
    return np.r_[today[release_slot:], tomorrow[:release_slot]]


def solve_schedule(
    load_forecast,
    pv_forecast,
    horizon_prices,
    e_start,
    base_plan=None,
    terminal_mode="fixed",
    terminal_band=600.0,
    terminal_penalty=None,
):
    model = pulp.LpProblem("Q3_asymmetric_settlement_mpc", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("buy", range(N), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(N), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(N), 0, Q_MAX)
    spill = pulp.LpVariable.dicts("spill", range(N), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(N + 1), E_MIN, E_MAX)
    model += energy[0] == float(e_start)
    terminal_terms = []
    if terminal_mode == "fixed":
        model += energy[N] == E_INITIAL
    elif terminal_mode == "band":
        model += energy[N] >= E_INITIAL - float(terminal_band)
        model += energy[N] <= E_INITIAL + float(terminal_band)
    elif terminal_mode == "soft":
        terminal_above = pulp.LpVariable("terminal_above", lowBound=0)
        terminal_below = pulp.LpVariable("terminal_below", lowBound=0)
        model += energy[N] - E_INITIAL == terminal_above - terminal_below
        penalty = (
            float(np.mean(horizon_prices))
            if terminal_penalty is None
            else float(terminal_penalty)
        )
        terminal_terms.append(penalty * (terminal_above + terminal_below))
    else:
        raise ValueError(f"未知终端SOC策略: {terminal_mode}")

    for k in range(N):
        model += (
            buy[k] + float(pv_forecast[k]) + discharge[k]
            == float(load_forecast[k]) + charge[k] + spill[k]
        )
        model += energy[k + 1] == energy[k] + ETA_C * charge[k] - discharge[k] / ETA_D

    cost_terms = []
    if base_plan is None:
        cost_terms = [float(horizon_prices[k]) * buy[k] for k in range(N)]
    else:
        for k in range(N):
            price = float(horizon_prices[k])
            if k < len(base_plan):
                increase = pulp.LpVariable(f"increase_{k}", lowBound=0)
                decrease = pulp.LpVariable(f"decrease_{k}", lowBound=0)
                model += buy[k] - float(base_plan[k]) == increase - decrease
                cost_terms.append(
                    price * (float(base_plan[k]) + 1.5 * increase - 0.5 * decrease)
                )
            else:
                cost_terms.append(price * buy[k])

    model += (
        pulp.lpSum(cost_terms)
        + pulp.lpSum(terminal_terms)
        + 1e-8 * pulp.lpSum(charge[k] + discharge[k] + spill[k] for k in range(N))
    )
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"问题三滚动优化失败: {pulp.LpStatus[status]}")
    return {
        "buy": np.asarray([pulp.value(buy[k]) for k in range(N)]),
        "charge": np.asarray([pulp.value(charge[k]) for k in range(N)]),
        "discharge": np.asarray([pulp.value(discharge[k]) for k in range(N)]),
        "spill": np.asarray([pulp.value(spill[k]) for k in range(N)]),
        "energy": np.asarray([pulp.value(energy[k]) for k in range(N + 1)]),
    }


def execute_block(buy, actual_load, actual_pv, e_start):
    length = len(buy)
    charge = np.zeros(length)
    discharge = np.zeros(length)
    emergency = np.zeros(length)
    spill = np.zeros(length)
    energy = np.zeros(length + 1)
    energy[0] = e_start
    current = float(e_start)
    for k in range(length):
        deficit = float(actual_load[k] - actual_pv[k] - buy[k])
        if deficit >= 0:
            discharge[k] = min(deficit, Q_MAX, max((current - E_MIN) * ETA_D, 0.0))
            emergency[k] = deficit - discharge[k]
        else:
            surplus = -deficit
            charge[k] = min(surplus, Q_MAX, max((E_MAX - current) / ETA_C, 0.0))
            spill[k] = surplus - charge[k]
        current += ETA_C * charge[k] - discharge[k] / ETA_D
        energy[k + 1] = current
    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "spill": spill,
        "energy": energy,
    }


def settlement(prices, plan, adjusted, emergency):
    increase = np.maximum(adjusted - plan, 0.0)
    decrease = np.maximum(plan - adjusted, 0.0)
    plan_cost = float(np.dot(prices, plan))
    net_adjustment = float(np.dot(prices, 1.5 * increase - 0.5 * decrease))
    retained = float(np.dot(prices, np.minimum(plan, adjusted)))
    cancellation = float(np.dot(0.5 * prices, decrease))
    incremental = float(np.dot(1.5 * prices, increase))
    scheduled_cost = retained + cancellation + incremental
    emergency_cost = float(np.dot(5.0 * prices, emergency))
    if abs((plan_cost + net_adjustment) - scheduled_cost) > 1e-5:
        raise AssertionError("调整购电费用分解不一致")
    return {
        "plan_cost": plan_cost,
        "net_adjustment_cost": net_adjustment,
        "retained_plan_cost": retained,
        "cancellation_penalty": cancellation,
        "incremental_purchase_cost": incremental,
        "scheduled_cost": scheduled_cost,
        "emergency_cost": emergency_cost,
        "total_cost": scheduled_cost + emergency_cost,
        "increase": increase,
        "decrease": decrease,
    }


def simulate_strategy(
    name,
    prices,
    prior_load,
    dates,
    loads,
    pvs,
    pv_forecasts,
    pv_margins,
    release_slots,
    robust,
    limit_days=None,
    terminal_mode="fixed",
    terminal_band=600.0,
    terminal_penalty=None,
    verbose=True,
):
    total_days = len(dates) if limit_days is None else min(len(dates), limit_days)
    results = []
    e_start = E_INITIAL
    max_balance = 0.0
    max_storage = 0.0

    for day in range(total_days):
        daily_plan = None
        adjusted = np.zeros(N)
        charge = np.zeros(N)
        discharge = np.zeros(N)
        emergency = np.zeros(N)
        spill = np.zeros(N)
        energy = np.zeros(N + 1)
        energy[0] = e_start
        current = float(e_start)
        releases = tuple(release_slots)

        for release_position, release_slot in enumerate(releases):
            release_index = RELEASE_SLOTS.index(release_slot)
            load_horizon = load_forecast_horizon(loads, day, release_slot, prior_load)
            pv_horizon = pv_forecasts[day, release_index].copy()
            if robust:
                pv_horizon = np.maximum(pv_horizon - pv_margins[day, release_index], 0.0)
            horizon_prices = prices[(release_slot + np.arange(N)) % N]
            base = None if release_slot == 0 else daily_plan[release_slot:]
            schedule = solve_schedule(
                load_horizon,
                pv_horizon,
                horizon_prices,
                current,
                base,
                terminal_mode=terminal_mode,
                terminal_band=terminal_band,
                terminal_penalty=terminal_penalty,
            )
            if release_slot == 0:
                daily_plan = schedule["buy"].copy()

            next_release = releases[release_position + 1] if release_position + 1 < len(releases) else N
            block_length = next_release - release_slot
            committed = schedule["buy"][:block_length]
            stop = release_slot + block_length
            adjusted[release_slot:stop] = committed
            actual = execute_block(
                committed,
                loads[day, release_slot:stop],
                pvs[day, release_slot:stop],
                current,
            )
            charge[release_slot:stop] = actual["charge"]
            discharge[release_slot:stop] = actual["discharge"]
            emergency[release_slot:stop] = actual["emergency"]
            spill[release_slot:stop] = actual["spill"]
            energy[release_slot : stop + 1] = actual["energy"]
            current = float(actual["energy"][-1])

        balance = adjusted + pvs[day] + discharge + emergency - loads[day] - charge - spill
        storage = energy[1:] - energy[:-1] - ETA_C * charge + discharge / ETA_D
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_storage = max(max_storage, float(np.max(np.abs(storage))))
        if np.min(energy) < E_MIN - 1e-4 or np.max(energy) > E_MAX + 1e-4:
            raise AssertionError(f"{name}第{day + 1}天SOC越界")
        if np.any((charge > TOL) & (discharge > TOL)):
            raise AssertionError(f"{name}第{day + 1}天同时充放电")

        costs = settlement(prices, daily_plan, adjusted, emergency)
        results.append(
            {
                "date": dates[day],
                "plan": daily_plan,
                "adjusted": adjusted,
                "charge": charge,
                "discharge": discharge,
                "emergency": emergency,
                "spill": spill,
                "energy": energy,
                **costs,
            }
        )
        e_start = current
        if verbose and ((day + 1) % 30 == 0 or day + 1 == total_days):
            print(f"{name}: {day + 1}/{total_days}")

    return results, {"balance": max_balance, "storage": max_storage}


def summarize(results):
    report = results[REPORT_START:] if len(results) > REPORT_START else results
    return {
        "days": len(report),
        "plan_kwh": sum(float(np.sum(r["plan"])) for r in report),
        "adjusted_kwh": sum(float(np.sum(r["adjusted"])) for r in report),
        "increase_kwh": sum(float(np.sum(r["increase"])) for r in report),
        "decrease_kwh": sum(float(np.sum(r["decrease"])) for r in report),
        "emergency_kwh": sum(float(np.sum(r["emergency"])) for r in report),
        "plan_cost": sum(r["plan_cost"] for r in report),
        "net_adjustment_cost": sum(r["net_adjustment_cost"] for r in report),
        "scheduled_cost": sum(r["scheduled_cost"] for r in report),
        "emergency_cost": sum(r["emergency_cost"] for r in report),
        "total_cost": sum(r["total_cost"] for r in report),
        "emergency_days": sum(float(np.sum(r["emergency"])) > 0.01 for r in report),
        "adjustment_days": sum(float(np.sum(np.abs(r["adjusted"] - r["plan"]))) > 0.01 for r in report),
        "end_energy": float(report[-1]["energy"][-1]) if report else np.nan,
    }


def source_order(values):
    values = list(values)
    return [*values[1:], values[0]]


def interval_label(start, stop):
    def hhmm(minutes):
        minutes %= 1440
        return f"{minutes // 60:02d}:{minutes % 60:02d}"

    suffix = "+1" if stop * 10 >= 1440 else ""
    return f"{hhmm(start * 10)}-{hhmm(stop * 10)}{suffix}"


def write_output(prices, dates, results):
    report = results[REPORT_START:]
    report_dates = dates[REPORT_START:]
    copyfile(TEMPLATE, OUTPUT)
    wb = openpyxl.load_workbook(OUTPUT)

    for sheet_name, key in (("计划购电量", "plan"), ("调整购电量", "adjusted")):
        ws = wb[sheet_name]
        for i, result in enumerate(report, start=2):
            for col, value in enumerate(source_order(result[key]), start=2):
                ws.cell(i, col, round(float(value), 4))
            ws.cell(i, 146, round(float(np.sum(result[key])), 4))
            cost = result["plan_cost"] if key == "plan" else result["scheduled_cost"]
            ws.cell(i, 147, round(float(cost), 4))

    ws = wb["充放电量"]
    styles = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    ws.delete_rows(2, ws.max_row)
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for date_value, result in zip(report_dates, report):
        for block in range(6):
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(styles[block][col - 1])
            start, stop = block * 24, (block + 1) * 24
            ws.cell(row, 1, date_value if block == 0 else None)
            ws.cell(row, 2, labels[block])
            ws.cell(row, 3, round(float(np.sum(result["charge"][start:stop])), 4))
            ws.cell(row, 4, round(float(np.sum(result["discharge"][start:stop])), 4))
            if block == 0:
                ws.cell(row, 5, time(0, 0))
                ws.cell(row, 6, round(float(result["energy"][0]), 4))
            elif block == 1:
                ws.cell(row, 5, "24:00")
                ws.cell(row, 6, round(float(result["energy"][-1]), 4))
            row += 1

    ws = wb["紧急购电量"]
    styles = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    ws.delete_rows(2, ws.max_row)
    row = 2
    for date_value, result in zip(report_dates, report):
        events = []
        t = 0
        while t < N:
            if result["emergency"][t] < EXPORT_TOL:
                t += 1
                continue
            start = t
            while t < N and result["emergency"][t] >= EXPORT_TOL:
                t += 1
            events.append((start, t))
        for event_index, (start, stop) in enumerate(events):
            style = styles[min(event_index, len(styles) - 1)]
            for col in range(1, 4):
                ws.cell(row, col)._style = copy(style[col - 1])
            ws.cell(row, 1, date_value if event_index == 0 else None)
            ws.cell(row, 2, interval_label(start, stop))
            ws.cell(row, 3, round(float(np.sum(result["emergency"][start:stop])), 4))
            row += 1

    wb.save(OUTPUT)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit-days", type=int, default=None)
    parser.add_argument("--selected-only", action="store_true")
    args = parser.parse_args()

    prices, prior_load, _, dates, loads, pvs = q2_final.load_inputs()
    pv_forecasts, pv_margins = load_pv_forecasts(dates, pvs)
    strategies = [
        ("static_robust", (0,), True),
        ("robust_06", (0, 36), True),
        ("robust_0612", (0, 36, 72), True),
        ("nominal_mpc", RELEASE_SLOTS, False),
        ("robust_mpc", RELEASE_SLOTS, True),
    ]
    if args.selected_only:
        strategies = [strategies[-1]]

    outputs = {}
    checks = {}
    for name, release_slots, robust in strategies:
        outputs[name], checks[name] = simulate_strategy(
            name,
            prices,
            prior_load,
            dates,
            loads,
            pvs,
            pv_forecasts,
            pv_margins,
            release_slots,
            robust,
            args.limit_days,
        )

    summaries = {name: summarize(result) for name, result in outputs.items()}
    if args.limit_days is None and "robust_mpc" in outputs:
        write_output(prices, dates, outputs["robust_mpc"])

    lines = ["Q3 FINAL MODEL METRICS"]
    for name, values in summaries.items():
        lines.append(f"[{name}]")
        lines.extend(f"{key}={value}" for key, value in values.items())
        lines.append(f"max_balance_residual={checks[name]['balance']}")
        lines.append(f"max_storage_residual={checks[name]['storage']}")
    if "static_robust" in summaries and "robust_mpc" in summaries:
        static = summaries["static_robust"]
        robust = summaries["robust_mpc"]
        lines.append("[information_value]")
        lines.append(f"total_cost_saving={static['total_cost'] - robust['total_cost']}")
        lines.append(f"emergency_kwh_reduction={static['emergency_kwh'] - robust['emergency_kwh']}")
    if all(name in summaries for name in ("static_robust", "robust_06", "robust_0612", "robust_mpc")):
        stage_names = ("static_robust", "robust_06", "robust_0612", "robust_mpc")
        release_names = ("06", "12", "18")
        lines.append("[marginal_update_value]")
        for before, after, release_name in zip(stage_names, stage_names[1:], release_names):
            lines.append(
                f"update_{release_name}_cost_saving="
                f"{summaries[before]['total_cost'] - summaries[after]['total_cost']}"
            )
            lines.append(
                f"update_{release_name}_emergency_reduction="
                f"{summaries[before]['emergency_kwh'] - summaries[after]['emergency_kwh']}"
            )
    text = "\n".join(lines)
    print(text)
    if args.limit_days is None:
        METRICS.write_text(text + "\n", encoding="utf-8")
        print(f"result={OUTPUT}")
        print(f"metrics={METRICS}")


if __name__ == "__main__":
    main()

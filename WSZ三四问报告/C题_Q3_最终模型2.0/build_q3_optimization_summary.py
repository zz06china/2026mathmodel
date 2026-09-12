from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "第三问优化与敏感性分析汇总.xlsx"


def parse_records(path):
    records = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if "=" not in line or "," not in line:
            continue
        record = {}
        for item in line.split(","):
            if "=" not in item:
                continue
            key, value = item.split("=", 1)
            record[key] = value
        if "total_cost" in record:
            records.append(record)
    return records


def number(record, key):
    return float(record[key])


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    for cell in sheet[1]:
        cell.font = Font(name="宋体", size=10.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium, bottom=thin)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for cell in sheet[sheet.max_row]:
        cell.border = Border(bottom=medium)
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = 20
    sheet.column_dimensions["A"].width = 27
    sheet.freeze_panes = "A2"


def main():
    uniform = parse_records(HERE / "q3_sensitivity.txt")[:5]
    staged = parse_records(HERE / "q3_stage_quantile.txt")
    selected = parse_records(HERE / "q3_selected_sensitivity.txt")

    headers = [
        "参数方案",
        "调整后购电费/万元",
        "紧急购电费/万元",
        "总费用/万元",
        "紧急购电量/kWh",
        "紧急购电天数",
    ]
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    stage_rows = []
    for record in staged:
        stage_rows.append(
            [
                f"q0=0.80，q调整={record['q_adjust']}",
                number(record, "scheduled_cost") / 10000,
                number(record, "emergency_cost") / 10000,
                number(record, "total_cost") / 10000,
                number(record, "emergency_kwh"),
                int(record["emergency_days"]),
            ]
        )
    baseline_uniform = uniform[2]
    stage_rows.append(
        [
            "q0=0.80，q调整=0.80",
            number(baseline_uniform, "scheduled_cost") / 10000,
            number(baseline_uniform, "emergency_cost") / 10000,
            number(baseline_uniform, "total_cost") / 10000,
            number(baseline_uniform, "emergency_kwh"),
            int(baseline_uniform["emergency_days"]),
        ]
    )
    add_sheet(workbook, "分阶段分位数", headers, stage_rows)

    uniform_rows = [
        [
            f"统一q={record['q']}",
            number(record, "scheduled_cost") / 10000,
            number(record, "emergency_cost") / 10000,
            number(record, "total_cost") / 10000,
            number(record, "emergency_kwh"),
            int(record["emergency_days"]),
        ]
        for record in uniform
    ]
    add_sheet(workbook, "统一分位数对照", headers, uniform_rows)

    final_baseline = {
        "scheduled_cost": 13023742.8622999,
        "emergency_cost": 639592.6831789727,
        "total_cost": 13663335.545478873,
        "emergency_kwh": 112555.17931659606,
        "emergency_days": 263,
    }
    window_records = {record["case"]: record for record in selected if record["case"].startswith("窗口")}
    window_rows = []
    for label in ("窗口14天",):
        record = window_records[label]
        window_rows.append(
            [
                label,
                number(record, "scheduled_cost") / 10000,
                number(record, "emergency_cost") / 10000,
                number(record, "total_cost") / 10000,
                number(record, "emergency_kwh"),
                int(record["emergency_days"]),
            ]
        )
    window_rows.append(
        [
            "窗口28天（采用）",
            final_baseline["scheduled_cost"] / 10000,
            final_baseline["emergency_cost"] / 10000,
            final_baseline["total_cost"] / 10000,
            final_baseline["emergency_kwh"],
            final_baseline["emergency_days"],
        ]
    )
    record = window_records["窗口56天"]
    window_rows.append(
        [
            "窗口56天",
            number(record, "scheduled_cost") / 10000,
            number(record, "emergency_cost") / 10000,
            number(record, "total_cost") / 10000,
            number(record, "emergency_kwh"),
            int(record["emergency_days"]),
        ]
    )
    add_sheet(workbook, "历史窗口", headers, window_rows)

    terminal_records = {record["case"]: record for record in selected if record["case"].startswith("终端")}
    terminal_rows = [
        [
            "固定6000 kWh（采用）",
            final_baseline["scheduled_cost"] / 10000,
            final_baseline["emergency_cost"] / 10000,
            final_baseline["total_cost"] / 10000,
            final_baseline["emergency_kwh"],
            final_baseline["emergency_days"],
        ]
    ]
    for key, label in (
        ("终端区间5400-6600 kWh", "允许5400-6600 kWh"),
        ("终端偏离软惩罚", "偏离6000 kWh软惩罚"),
    ):
        record = terminal_records[key]
        terminal_rows.append(
            [
                label,
                number(record, "scheduled_cost") / 10000,
                number(record, "emergency_cost") / 10000,
                number(record, "total_cost") / 10000,
                number(record, "emergency_kwh"),
                int(record["emergency_days"]),
            ]
        )
    add_sheet(workbook, "终端SOC", headers, terminal_rows)
    workbook.save(OUTPUT)
    print(f"saved={OUTPUT}")


if __name__ == "__main__":
    main()

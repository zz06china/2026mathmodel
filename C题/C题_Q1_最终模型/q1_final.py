from pathlib import Path
from shutil import copyfile

import openpyxl
import pulp


ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = ROOT / "C题" / "附件" / "附件1.xlsx"
TEMPLATE_FILE = ROOT / "C题" / "附件" / "附件5" / "result1.xlsx"
OUTPUT_FILE = Path(__file__).resolve().parent / "result1_final.xlsx"

DT = 1.0 / 6.0
N = 144
ETA_C = 0.9
ETA_D = 0.9
E_INITIAL = 6000.0
E_MIN = 1200.0
E_MAX = 10800.0
P_STORAGE_MAX = 5000.0
Q_STORAGE_MAX = P_STORAGE_MAX * DT
TOL = 1e-6


def time_label(value):
    if isinstance(value, str):
        return value.strip()
    return f"{value.hour}:{value.minute:02d}"


def load_data():
    workbook = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    sheet = workbook["Sheet1"]
    raw = list(sheet.iter_rows(min_row=2, values_only=True))
    if len(raw) != N:
        raise ValueError(f"附件1应包含{N}个时段，实际为{len(raw)}个")

    records = [
        {
            "source_index": i,
            "start": time_label(row[0]),
            "price": float(row[1]),
            "load": float(row[2]) * DT,
            "pv": float(row[3]) * DT,
        }
        for i, row in enumerate(raw)
    ]

    # 附件从0:10开始、以次日0:00结束。题目要求在0:00开始调度，
    # 因每天数据相同，将末行0:00+1旋转到首位，作为0:00-0:10时段。
    return [records[-1], *records[:-1]], records


def build_and_solve(records):
    model = pulp.LpProblem("Q1_microgrid_dispatch", pulp.LpMinimize)

    buy = pulp.LpVariable.dicts("buy", range(N), lowBound=0)
    charge = pulp.LpVariable.dicts(
        "charge", range(N), lowBound=0, upBound=Q_STORAGE_MAX
    )
    discharge = pulp.LpVariable.dicts(
        "discharge", range(N), lowBound=0, upBound=Q_STORAGE_MAX
    )
    curtail = {
        t: pulp.LpVariable("curtail_%d" % t, lowBound=0, upBound=records[t]["pv"])
        for t in range(N)
    }
    energy = pulp.LpVariable.dicts("energy", range(N + 1), E_MIN, E_MAX)

    model += energy[0] == E_INITIAL, "initial_energy"
    model += energy[N] == E_INITIAL, "terminal_energy"

    for t, row in enumerate(records):
        model += (
            buy[t] + row["pv"] + discharge[t]
            == row["load"] + charge[t] + curtail[t]
        ), f"balance_{t}"
        model += (
            energy[t + 1]
            == energy[t] + ETA_C * charge[t] - discharge[t] / ETA_D
        ), f"storage_{t}"

    cost = pulp.lpSum(records[t]["price"] * buy[t] for t in range(N))
    model.setObjective(cost)
    solver = pulp.PULP_CBC_CMD(msg=False)
    status = model.solve(solver)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"第一阶段求解失败：{pulp.LpStatus[status]}")
    optimal_cost = pulp.value(cost)

    # 在保持最低购电费的前提下，选取吞吐量和弃光量最小的物理解。
    model += cost <= optimal_cost + 1e-5, "preserve_optimal_cost"
    model.setObjective(
        pulp.lpSum(charge[t] + discharge[t] + curtail[t] for t in range(N))
    )
    status = model.solve(solver)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"第二阶段求解失败：{pulp.LpStatus[status]}")

    result = {
        "buy": [pulp.value(buy[t]) for t in range(N)],
        "charge": [pulp.value(charge[t]) for t in range(N)],
        "discharge": [pulp.value(discharge[t]) for t in range(N)],
        "curtail": [pulp.value(curtail[t]) for t in range(N)],
        "energy": [pulp.value(energy[t]) for t in range(N + 1)],
    }
    result["cost"] = sum(records[t]["price"] * result["buy"][t] for t in range(N))
    result["optimal_cost_stage1"] = optimal_cost
    return result


def validate(records, result):
    residuals = [
        result["buy"][t]
        + records[t]["pv"]
        + result["discharge"][t]
        - records[t]["load"]
        - result["charge"][t]
        - result["curtail"][t]
        for t in range(N)
    ]
    storage_residuals = [
        result["energy"][t + 1]
        - result["energy"][t]
        - ETA_C * result["charge"][t]
        + result["discharge"][t] / ETA_D
        for t in range(N)
    ]
    simultaneous = [
        t
        for t in range(N)
        if result["charge"][t] > TOL and result["discharge"][t] > TOL
    ]
    checks = {
        "max_balance_error": max(abs(x) for x in residuals),
        "max_storage_error": max(abs(x) for x in storage_residuals),
        "min_energy": min(result["energy"]),
        "max_energy": max(result["energy"]),
        "terminal_error": abs(result["energy"][-1] - E_INITIAL),
        "max_charge": max(result["charge"]),
        "max_discharge": max(result["discharge"]),
        "simultaneous_periods": simultaneous,
    }
    # CBC导出的变量值保留有限位小数；1e-3 kWh远小于数据有效精度。
    if checks["max_balance_error"] > 1e-3:
        raise AssertionError("能量平衡误差过大")
    if checks["max_storage_error"] > 1e-3:
        raise AssertionError("储能递推误差过大")
    if checks["min_energy"] < E_MIN - 1e-4 or checks["max_energy"] > E_MAX + 1e-4:
        raise AssertionError("储电量越界")
    if checks["terminal_error"] > 1e-4:
        raise AssertionError("期末储电量不等于期初")
    if simultaneous:
        raise AssertionError(f"存在同时充放电时段：{simultaneous}")
    return checks


def baseline(records):
    buy = [max(row["load"] - row["pv"], 0.0) for row in records]
    cost = sum(row["price"] * buy[t] for t, row in enumerate(records))
    return sum(buy), cost


def source_order(values):
    # 自然日顺序为[原始143, 原始0, ..., 原始142]，模板顺序为原始0...143。
    return [*values[1:], values[0]]


def write_result(result):
    copyfile(TEMPLATE_FILE, OUTPUT_FILE)
    workbook = openpyxl.load_workbook(OUTPUT_FILE)
    plan_sheet = workbook["计划购电量"]
    for row, value in enumerate(source_order(result["buy"]), start=2):
        plan_sheet.cell(row=row, column=2, value=round(value, 4))

    storage_sheet = workbook["充放电量"]
    for block in range(6):
        start = block * 24
        stop = start + 24
        storage_sheet.cell(
            row=block + 2,
            column=2,
            value=round(sum(result["charge"][start:stop]), 4),
        )
        storage_sheet.cell(
            row=block + 2,
            column=3,
            value=round(sum(result["discharge"][start:stop]), 4),
        )
    storage_sheet.cell(row=2, column=5, value=round(result["energy"][0], 4))
    storage_sheet.cell(row=3, column=5, value=round(result["energy"][-1], 4))
    workbook.save(OUTPUT_FILE)


def main():
    records, _ = load_data()
    result = build_and_solve(records)
    checks = validate(records, result)
    base_buy, base_cost = baseline(records)
    write_result(result)

    by_start = {row["start"]: t for t, row in enumerate(records)}
    specified = ["10:00", "12:00", "14:00", "16:00", "18:00", "20:00"]

    print("求解状态: Optimal")
    print(f"第一阶段最低购电费: {result['optimal_cost_stage1']:.6f} 元")
    print(f"最终购电费: {result['cost']:.6f} 元")
    print(f"全天购电量: {sum(result['buy']):.6f} kWh")
    print(f"全天充电量: {sum(result['charge']):.6f} kWh")
    print(f"全天放电量: {sum(result['discharge']):.6f} kWh")
    print(f"全天弃光量: {sum(result['curtail']):.6f} kWh")
    print(f"无储能购电量: {base_buy:.6f} kWh")
    print(f"无储能购电费: {base_cost:.6f} 元")
    print(f"节省费用: {base_cost - result['cost']:.6f} 元")
    print(f"节费率: {(base_cost - result['cost']) / base_cost:.6%}")
    print("指定时段购电量:")
    for start in specified:
        print(f"  {start}: {result['buy'][by_start[start]]:.6f} kWh")
    print("4小时时段充放电量:")
    for block in range(6):
        start = block * 24
        stop = start + 24
        print(
            f"  {block * 4:02d}:00-{(block + 1) * 4:02d}:00: "
            f"充电 {sum(result['charge'][start:stop]):.6f}, "
            f"放电 {sum(result['discharge'][start:stop]):.6f} kWh"
        )
    print("校验:")
    for key, value in checks.items():
        print(f"  {key}: {value}")
    print(f"结果文件: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()

"""
C题 问题4：波动电价下的购电-储能联合优化（重新计算问题2与问题3）。

产出：
  - result4-2.xlsx：在附件4时变电价下重算问题2（0:00 日前计划 + 因果执行 + 紧急购电）。
  - result4-3.xlsx：在附件4时变电价下重算问题3（0/6/12/18 滚动 MPC + 调整购电计价）。

模型机制分别继承：
  - 问题2（q2_final.py）：自适应滞后融合负荷/光伏预报 + 48h 前瞻 + 报童安全裕量；
  - 问题3（problem3.py）：附件3光伏预报锚点插值 + 承诺时刻残差 + 调整购电 1.5×/50% 计价。

唯一变化：p -> p_t（附件4 时变电价）。全部纯 LP（无 0-1 变量），充放电互斥由
往返效率 0.81 自动保证（同时充放电严格劣）。终端 SOC 维持日内循环（与问题3一致）。
"""

from __future__ import annotations

from copy import copy
from pathlib import Path
from shutil import copyfile

import numpy as np
import openpyxl
import pulp

BASE = Path(__file__).resolve().parent.parent  # C题/
ATTACHMENT1 = BASE / "附件" / "附件1.xlsx"
ATTACHMENT2 = BASE / "附件" / "附件2.xlsx"
ATTACHMENT3 = BASE / "附件" / "附件3.xlsx"
ATTACHMENT4 = BASE / "附件" / "附件4.xlsx"
TEMPLATE4_2 = BASE / "附件" / "附件5" / "result4-2.xlsx"
TEMPLATE4_3 = BASE / "附件" / "附件5" / "result4-3.xlsx"
OUTPUT4_2 = Path(__file__).resolve().parent / "result4-2.xlsx"
OUTPUT4_3 = Path(__file__).resolve().parent / "result4-3.xlsx"
SUMMARY4_2 = Path(__file__).resolve().parent / "q4_2_metrics.txt"
SUMMARY4_3 = Path(__file__).resolve().parent / "q4_3_metrics.txt"

DT = 1.0 / 6.0
N = 144
ETA_C = 0.9
ETA_D = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
Q_MAX = 5000.0 * DT
REPORT_START = 31
ALPHA_GRID = np.linspace(0.0, 1.0, 11)
RISK_QUANTILE = 0.80
RESIDUAL_WINDOW = 28
SOLVER = pulp.PULP_CBC_CMD(msg=False)
TOL = 1e-5
EXPORT_TOL = 5e-5
LEX_TOL = 1e-3  # 词典序第二阶段成本容差（元）
ISSUE_HOURS = [0, 6, 12, 18]


# ---------------------------------------------------------------------------
# 时间轴与数据加载
# ---------------------------------------------------------------------------
def rotate_to_natural_day(values):
    values = list(values)
    if len(values) != N:
        raise ValueError(f"Expected {N} intervals, got {len(values)}")
    return np.asarray([values[-1], *values[:-1]], dtype=float)


def source_order(values):
    values = list(values)
    return [*values[1:], values[0]]


def load_inputs():
    wb1 = openpyxl.load_workbook(ATTACHMENT1, data_only=True, read_only=True)
    rows = list(wb1["Sheet1"].iter_rows(min_row=2, values_only=True))
    wb1.close()
    prices = rotate_to_natural_day(float(row[1]) for row in rows)
    prior_load = rotate_to_natural_day(float(row[2]) * DT for row in rows)
    prior_pv = rotate_to_natural_day(float(row[3]) * DT for row in rows)

    wb2 = openpyxl.load_workbook(ATTACHMENT2, data_only=True, read_only=True)
    load_sheet, pv_sheet = wb2.worksheets[:2]

    def read_sheet(sheet):
        data, dates = [], []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            dates.append(row[0])
            data.append(rotate_to_natural_day(float(v) * DT for v in row[1 : N + 1]))
        return dates, np.asarray(data)

    dates, loads = read_sheet(load_sheet)
    pv_dates, pvs = read_sheet(pv_sheet)
    wb2.close()
    return prices, prior_load, prior_pv, dates, loads, pvs


def load_forecasts():
    """附件3：shape (365, 4, 24)，[day][issue][预报1..24小时]."""
    wb = openpyxl.load_workbook(ATTACHMENT3, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == 365 * 4, f"附件3行数={len(rows)}"
    forecasts = np.zeros((365, 4, 24))

    def issue_index(label):
        hour = int(str(label).strip().split(":")[0])
        return {0: 0, 6: 1, 12: 2, 18: 3}[hour]

    day = -1
    for r in rows:
        if r[0] is not None and str(r[0]).strip() != "":
            day += 1
        idx = issue_index(r[1])
        forecasts[day, idx] = np.array([float(v) for v in r[2:26]])
    return forecasts


def load_prices4():
    """附件4：365×144 时变电价（元/kWh），按自然日排列，与附件2日期对齐."""
    wb = openpyxl.load_workbook(ATTACHMENT4, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        data.append(rotate_to_natural_day(float(v) for v in row[1 : N + 1]))
    wb.close()
    arr = np.asarray(data)
    assert arr.shape == (365, N), f"附件4维度异常: {arr.shape}"
    return arr


# ---------------------------------------------------------------------------
# 预报：负载（自适应滞后融合）与光伏（问题2滞后融合 / 问题3附件3锚点插值）
# ---------------------------------------------------------------------------
def select_alpha(data, day, horizon, prior):
    """只用 day 0:00 之前可观测的数据选择滞后融合系数."""
    max_validation_day = day - 1
    start = max(7, max_validation_day - 55)
    candidates = []
    for alpha in ALPHA_GRID:
        errors = []
        for k in range(start, max_validation_day + 1):
            if horizon == 0:
                pred = alpha * data[k - 1] + (1.0 - alpha) * data[k - 7]
            else:
                if k < 7:
                    continue
                pred = alpha * data[k - 2] + (1.0 - alpha) * data[k - 7]
            errors.append(np.mean(np.abs(data[k] - pred)))
        candidates.append(np.mean(errors) if errors else np.inf)
    if not np.isfinite(min(candidates)):
        return 0.0
    return float(ALPHA_GRID[int(np.argmin(candidates))])


def point_forecast(data, day, horizon, prior):
    """问题2式滞后融合点预测；horizon=0 为当日、1 为次日（只用滞后观测）."""
    if day < 7:
        return prior.copy(), 0.0
    alpha = select_alpha(data, day, horizon, prior)
    if horizon == 0:
        forecast = alpha * data[day - 1] + (1.0 - alpha) * data[day - 7]
    else:
        forecast = alpha * data[day - 1] + (1.0 - alpha) * data[day - 6]
    return np.maximum(forecast, 0.0), alpha


def load_forecast_10min(loads, day, prior_load):
    """日负荷10分钟预测（问题3用，只在0:00做一次，等价 point_forecast horizon=0）."""
    return point_forecast(loads, day, 0, prior_load)[0]


def hourly_forecast_to_10min(hourly, issue_hour, anchor=0.0):
    """把24个整点预报值插值到144个10分钟时段。

    hourly[k] = 光伏在 (issue_hour + k + 1) 点的预报值，k=0..23。
    发布时刻（issue_hour）用发布前最后一个实际光伏值 anchor 作零时刻锚点
    （单位同 marker，kWh/10min），避免当前小时借下一整点预报造成水平段。
    """
    marker = np.zeros(25)  # marker[h] = 小时 h 的光伏能量（kWh/10min = kW * DT）
    for h in range(1, 25):
        off = h - issue_hour
        if 1 <= off <= 24:
            marker[h] = hourly[off - 1] * DT
    if issue_hour > 0:
        marker[issue_hour] = anchor
    minutes = (np.arange(N) * 10 + 5) / 60.0  # 每时段中点在“小时”单位的坐标
    return np.interp(minutes, np.arange(25), marker)


def residual_margin(residual_history, window=RESIDUAL_WINDOW, q=RISK_QUANTILE):
    """净负荷预测残差的 q 分位修正（报童安全库存）。"""
    if not residual_history:
        return np.zeros(N)
    sample = np.asarray(residual_history[-window:])
    return np.quantile(sample, q, axis=0)


# ---------------------------------------------------------------------------
# LP 求解
# ---------------------------------------------------------------------------
def _finish_two_stage(model, cost, smoothing_vars):
    """第一阶段最小化 cost，第二阶段在容差内最小化吞吐量/弃电以去退化."""
    model.setObjective(cost)
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"第一阶段求解失败: {pulp.LpStatus[status]}")
    optimum = pulp.value(cost)
    model += cost <= optimum + LEX_TOL
    model.setObjective(pulp.lpSum(smoothing_vars))
    status = model.solve(SOLVER)
    if pulp.LpStatus[status] != "Optimal":
        raise RuntimeError(f"第二阶段求解失败: {pulp.LpStatus[status]}")


def solve_horizon(net_targets, price_list, e_start, terminal_target=E_INITIAL):
    """问题2式滚动计划：horizon 天（各 144 时段），price_list 为每天电价."""
    horizon = len(net_targets)
    total_slots = horizon * N
    model = pulp.LpProblem("Q4_rolling_plan", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("buy", range(total_slots), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(total_slots), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(total_slots), 0, Q_MAX)
    surplus = pulp.LpVariable.dicts("surplus", range(total_slots), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(total_slots + 1), E_MIN, E_MAX)
    model += energy[0] == e_start
    model += energy[total_slots] == terminal_target

    for j in range(horizon):
        for t in range(N):
            k = j * N + t
            model += buy[k] + discharge[k] == float(net_targets[j][t]) + charge[k] + surplus[k]
            model += energy[k + 1] == energy[k] + ETA_C * charge[k] - discharge[k] / ETA_D

    cost = pulp.lpSum(
        price_list[j][t] * buy[j * N + t] for j in range(horizon) for t in range(N)
    )
    _finish_two_stage(model, cost, [charge[k] + discharge[k] + surplus[k] for k in range(total_slots)])

    first = range(N)
    return {
        "buy": np.array([pulp.value(buy[t]) for t in first]),
        "charge": np.array([pulp.value(charge[t]) for t in first]),
        "discharge": np.array([pulp.value(discharge[t]) for t in first]),
        "planned_surplus": np.array([pulp.value(surplus[t]) for t in first]),
        "energy": np.array([pulp.value(energy[t]) for t in range(N + 1)]),
    }


def solve_dayahead(net_targets, prices, e_start, terminal=E_INITIAL):
    """0:00 日前计划：min Σ p B，得到计划购电量 B_t（144时段）."""
    model = pulp.LpProblem("Q4_dayahead", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("buy", range(N), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(N), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(N), 0, Q_MAX)
    surplus = pulp.LpVariable.dicts("surplus", range(N), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(N + 1), E_MIN, E_MAX)

    model += energy[0] == e_start
    model += energy[N] == terminal
    for t in range(N):
        model += buy[t] + discharge[t] == float(net_targets[t]) + charge[t] + surplus[t]
        model += energy[t + 1] == energy[t] + ETA_C * charge[t] - discharge[t] / ETA_D

    cost = pulp.lpSum(prices[t] * buy[t] for t in range(N))
    _finish_two_stage(model, cost, [charge[t] + discharge[t] + surplus[t] for t in range(N)])
    return {
        "buy": np.array([pulp.value(buy[t]) for t in range(N)]),
        "charge": np.array([pulp.value(charge[t]) for t in range(N)]),
        "discharge": np.array([pulp.value(discharge[t]) for t in range(N)]),
        "surplus": np.array([pulp.value(surplus[t]) for t in range(N)]),
        "energy": np.array([pulp.value(energy[t]) for t in range(N + 1)]),
    }


def solve_adjust(B_full, net_input, prices, e_start, start_slot, terminal=E_INITIAL):
    """在 start_slot 时刻重优化剩余时段，含相对计划 B 的调整成本。

    目标：min Σ [p A + 0.5 p (Δ_up + Δ_down)]，约束 A - Δ_up + Δ_down = B。
    上调价 1.5p、下调违约 0.5p，对 A 凸分段线性，纯 LP，最优解自动 Δ_up·Δ_down=0。
    """
    S = start_slot
    M = N - S
    model = pulp.LpProblem(f"Q4_adjust_{S}", pulp.LpMinimize)
    buy = pulp.LpVariable.dicts("a", range(M), lowBound=0)
    charge = pulp.LpVariable.dicts("charge", range(M), 0, Q_MAX)
    discharge = pulp.LpVariable.dicts("discharge", range(M), 0, Q_MAX)
    surplus = pulp.LpVariable.dicts("surplus", range(M), lowBound=0)
    dup = pulp.LpVariable.dicts("dup", range(M), lowBound=0)
    ddown = pulp.LpVariable.dicts("ddown", range(M), lowBound=0)
    energy = pulp.LpVariable.dicts("energy", range(M + 1), E_MIN, E_MAX)

    model += energy[0] == e_start
    model += energy[M] == terminal
    for i in range(M):
        t = S + i
        model += buy[i] - dup[i] + ddown[i] == float(B_full[t])
        model += (
            buy[i] + discharge[i]
            == float(net_input[t]) + charge[i] + surplus[i]
        )
        model += energy[i + 1] == energy[i] + ETA_C * charge[i] - discharge[i] / ETA_D

    cost = pulp.lpSum(
        prices[S + i] * (buy[i] + 0.5 * (dup[i] + ddown[i])) for i in range(M)
    )
    _finish_two_stage(
        model, cost, [charge[i] + discharge[i] + surplus[i] for i in range(M)]
    )
    return np.array([pulp.value(buy[i]) for i in range(M)])


# ---------------------------------------------------------------------------
# 因果执行
# ---------------------------------------------------------------------------
def execute_causal_day(buy, load, pv, e_start):
    """问题2式因果执行：计划购电量 buy 锁定，电池按实际净负荷平衡，缺口才紧急购电."""
    energy = float(e_start)
    charge = np.zeros(N)
    discharge = np.zeros(N)
    emergency = np.zeros(N)
    surplus = np.zeros(N)
    trajectory = np.zeros(N + 1)
    trajectory[0] = energy
    for t in range(N):
        net_deficit = load[t] - pv[t] - buy[t]
        if net_deficit >= 0.0:
            discharge[t] = min(net_deficit, Q_MAX, max((energy - E_MIN) * ETA_D, 0.0))
            emergency[t] = net_deficit - discharge[t]
        else:
            available = -net_deficit
            charge[t] = min(available, Q_MAX, max((E_MAX - energy) / ETA_C, 0.0))
            surplus[t] = available - charge[t]
        energy += ETA_C * charge[t] - discharge[t] / ETA_D
        trajectory[t + 1] = energy
    return {
        "charge": charge,
        "discharge": discharge,
        "emergency": emergency,
        "surplus": surplus,
        "energy": trajectory,
    }


def execute_range(buy, load, pv, energy, start, stop, charge, discharge, emergency, surplus):
    """问题3式分段因果执行：在 [start, stop) 按实际观测平衡，缺口才紧急购电."""
    for i in range(start, stop):
        net_deficit = load[i] - pv[i] - buy[i]
        if net_deficit >= 0.0:
            discharge[i] = min(net_deficit, Q_MAX, max((energy - E_MIN) * ETA_D, 0.0))
            emergency[i] = net_deficit - discharge[i]
            charge[i] = 0.0
            surplus[i] = 0.0
        else:
            available = -net_deficit
            charge[i] = min(available, Q_MAX, max((E_MAX - energy) / ETA_C, 0.0))
            surplus[i] = available - charge[i]
            discharge[i] = 0.0
            emergency[i] = 0.0
        energy += ETA_C * charge[i] - discharge[i] / ETA_D
    return energy


def _build_energy_trajectory(e_start, charge, discharge):
    traj = np.zeros(N + 1)
    traj[0] = e_start
    for t in range(N):
        traj[t + 1] = traj[t] + ETA_C * charge[t] - discharge[t] / ETA_D
    return traj


# ---------------------------------------------------------------------------
# 校验
# ---------------------------------------------------------------------------
def validate_day(charge, discharge, energy, emergency, surplus, buy, load, pv):
    balance = buy + pv + discharge + emergency - load - charge - surplus
    storage = energy[1:] - energy[:-1] - ETA_C * charge + discharge / ETA_D
    simultaneous = np.flatnonzero((charge > TOL) & (discharge > TOL))
    return {
        "balance": float(np.max(np.abs(balance))),
        "storage": float(np.max(np.abs(storage))),
        "min_energy": float(np.min(energy)),
        "max_energy": float(np.max(energy)),
        "simultaneous": int(len(simultaneous)),
    }


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------
def interval_label(start, stop):
    def hhmm(minutes):
        minutes %= 1440
        return f"{minutes // 60:02d}:{minutes % 60:02d}"

    start_min = start * 10
    stop_min = stop * 10
    suffix = "+1" if stop_min >= 1440 else ""
    return f"{hhmm(start_min)}-{hhmm(stop_min)}{suffix}"


def write_output_4_2(dates, results):
    copyfile(TEMPLATE4_2, OUTPUT4_2)
    wb = openpyxl.load_workbook(OUTPUT4_2)

    # 计划购电量：B_t；最后两列 全天购电量 / 全天购电费（计划费）
    ws = wb["计划购电量"]
    for i, r in enumerate(results):
        row = i + 2
        for t, v in enumerate(source_order(r["buy"]), start=2):
            ws.cell(row, t, round(float(v), 4))
        ws.cell(row, 146, round(float(np.sum(r["buy"])), 4))
        ws.cell(row, 147, round(float(r["plan_cost"]), 4))

    # 充放电量 + 储电量（因果执行的实际值）
    ws = wb["充放电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    ws.delete_rows(2, ws.max_row)
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for date, r in zip(dates, results):
        for block in range(6):
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(style_rows[block][col - 1])
            start, stop = block * 24, (block + 1) * 24
            ws.cell(row, 1, date if block == 0 else None)
            ws.cell(row, 2, labels[block])
            ws.cell(row, 3, round(float(np.sum(r["charge"][start:stop])), 4))
            ws.cell(row, 4, round(float(np.sum(r["discharge"][start:stop])), 4))
            if block == 0:
                ws.cell(row, 5, "0:00")
                ws.cell(row, 6, round(float(r["energy"][0]), 4))
            elif block == 1:
                ws.cell(row, 5, "24:00")
                ws.cell(row, 6, round(float(r["energy"][-1]), 4))
            row += 1

    # 紧急购电量
    ws = wb["紧急购电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    ws.delete_rows(2, ws.max_row)
    row = 2
    for date, r in zip(dates, results):
        emergency = r["emergency"]
        starts, t = [], 0
        while t < N:
            if emergency[t] < EXPORT_TOL:
                t += 1
                continue
            start = t
            while t < N and emergency[t] >= EXPORT_TOL:
                t += 1
            starts.append((start, t))
        for event_index, (start, stop) in enumerate(starts):
            style = style_rows[min(event_index, len(style_rows) - 1)]
            for col in range(1, 4):
                ws.cell(row, col)._style = copy(style[col - 1])
            ws.cell(row, 1, date if event_index == 0 else None)
            ws.cell(row, 2, interval_label(start, stop))
            ws.cell(row, 3, round(float(np.sum(emergency[start:stop])), 4))
            row += 1

    wb.save(OUTPUT4_2)


def write_output_4_3(dates, results):
    copyfile(TEMPLATE4_3, OUTPUT4_3)
    wb = openpyxl.load_workbook(OUTPUT4_3)

    # 计划购电量：B_t；最后两列 全天购电量 / 全天购电费（计划费）
    ws = wb["计划购电量"]
    for i, r in enumerate(results):
        row = i + 2
        for t, v in enumerate(source_order(r["buy"]), start=2):
            ws.cell(row, t, round(float(v), 4))
        ws.cell(row, 146, round(float(np.sum(r["buy"])), 4))
        ws.cell(row, 147, round(float(r["plan_cost"]), 4))

    # 调整购电量：A_t；最后两列 全天调整购电量 / 全天购电费（调整后正常购电费）
    ws = wb["调整购电量"]
    for i, r in enumerate(results):
        row = i + 2
        for t, v in enumerate(source_order(r["adjusted"]), start=2):
            ws.cell(row, t, round(float(v), 4))
        ws.cell(row, 146, round(float(np.sum(r["adjusted"])), 4))
        ws.cell(row, 147, round(float(r["plan_cost"] + r["adjust_cost"]), 4))

    # 充放电量 + 储电量
    ws = wb["充放电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    ws.delete_rows(2, ws.max_row)
    labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    row = 2
    for date, r in zip(dates, results):
        for block in range(6):
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(style_rows[block][col - 1])
            start, stop = block * 24, (block + 1) * 24
            ws.cell(row, 1, date if block == 0 else None)
            ws.cell(row, 2, labels[block])
            ws.cell(row, 3, round(float(np.sum(r["charge"][start:stop])), 4))
            ws.cell(row, 4, round(float(np.sum(r["discharge"][start:stop])), 4))
            if block == 0:
                ws.cell(row, 5, "0:00")
                ws.cell(row, 6, round(float(r["energy"][0]), 4))
            elif block == 1:
                ws.cell(row, 5, "24:00")
                ws.cell(row, 6, round(float(r["energy"][-1]), 4))
            row += 1

    # 紧急购电量
    ws = wb["紧急购电量"]
    style_rows = [[copy(ws.cell(r, c)._style) for c in range(1, 4)] for r in range(2, 5)]
    ws.delete_rows(2, ws.max_row)
    row = 2
    for date, r in zip(dates, results):
        emergency = r["emergency"]
        starts, t = [], 0
        while t < N:
            if emergency[t] < EXPORT_TOL:
                t += 1
                continue
            start = t
            while t < N and emergency[t] >= EXPORT_TOL:
                t += 1
            starts.append((start, t))
        for event_index, (start, stop) in enumerate(starts):
            style = style_rows[min(event_index, len(style_rows) - 1)]
            for col in range(1, 4):
                ws.cell(row, col)._style = copy(style[col - 1])
            ws.cell(row, 1, date if event_index == 0 else None)
            ws.cell(row, 2, interval_label(start, stop))
            ws.cell(row, 3, round(float(np.sum(emergency[start:stop])), 4))
            row += 1

    wb.save(OUTPUT4_3)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def run_q2(prices4):
    """在时变电价下重算问题2：0:00 日前计划 + 因果执行 + 紧急购电."""
    _, prior_load, prior_pv, dates, loads, pvs = load_inputs()
    residuals = []
    all_results = []
    max_checks = {"balance": 0.0, "storage": 0.0, "min_energy": E_MAX, "max_energy": E_MIN}
    e_start = E_INITIAL

    for day in range(len(dates)):
        p_t = prices4[day]
        load0, _ = point_forecast(loads, day, 0, prior_load)
        pv0, _ = point_forecast(pvs, day, 0, prior_pv)
        net0 = load0 - pv0
        margin = residual_margin(residuals)
        targets = [net0 + margin]

        # 48h 前瞻：第二天价格沿用当天（避免年末放电，也不引入次日价格信息）
        load1, _ = point_forecast(loads, day, 1, prior_load)
        pv1, _ = point_forecast(pvs, day, 1, prior_pv)
        targets.append(load1 - pv1 + margin)

        plan = solve_horizon(targets, [p_t, p_t], e_start)
        actual = execute_causal_day(plan["buy"], loads[day], pvs[day], e_start)
        result = {
            "buy": plan["buy"],
            "charge": actual["charge"],
            "discharge": actual["discharge"],
            "energy": actual["energy"],
            "emergency": actual["emergency"],
            "surplus": actual["surplus"],
            "plan_cost": float(np.dot(p_t, plan["buy"])),
            "emergency_cost": float(np.dot(5.0 * p_t, actual["emergency"])),
        }
        checks = validate_day(
            actual["charge"], actual["discharge"], actual["energy"],
            actual["emergency"], actual["surplus"], plan["buy"], loads[day], pvs[day],
        )
        max_checks["balance"] = max(max_checks["balance"], checks["balance"])
        max_checks["storage"] = max(max_checks["storage"], checks["storage"])
        max_checks["min_energy"] = min(max_checks["min_energy"], checks["min_energy"])
        max_checks["max_energy"] = max(max_checks["max_energy"], checks["max_energy"])

        all_results.append(result)
        residuals.append((loads[day] - pvs[day]) - net0)
        e_start = float(actual["energy"][-1])

    report = all_results[REPORT_START:]
    report_dates = dates[REPORT_START:]

    plan_kwh = sum(np.sum(r["buy"]) for r in report)
    emergency_kwh = sum(np.sum(r["emergency"]) for r in report)
    plan_cost = sum(r["plan_cost"] for r in report)
    emergency_cost = sum(r["emergency_cost"] for r in report)
    total_cost = plan_cost + emergency_cost
    metrics = {
        "report_days": len(report),
        "planned_purchase": plan_kwh,
        "emergency_purchase": emergency_kwh,
        "planned_cost": plan_cost,
        "emergency_cost": emergency_cost,
        "total_cost": total_cost,
        "emergency_days": sum(np.sum(r["emergency"]) > 0.01 for r in report),
        "ending_energy": float(report[-1]["energy"][-1]),
        "max_balance_residual": max_checks["balance"],
        "max_storage_residual": max_checks["storage"],
        "global_min_energy": max_checks["min_energy"],
        "global_max_energy": max_checks["max_energy"],
        "simultaneous": sum(int(np.sum((r["charge"] > TOL) & (r["discharge"] > TOL))) for r in report),
    }
    write_output_4_2(report_dates, report)
    return metrics


def run_q3(prices4, adjust_issues=(6, 12, 18)):
    """在时变电价下重算问题3：0/6/12/18 滚动 MPC + 调整购电计价."""
    _, prior_load, prior_pv, dates, loads, pvs = load_inputs()
    forecasts = load_forecasts()
    adjust_issues = sorted(adjust_issues)
    boundaries = [h * 6 for h in adjust_issues]
    first_boundary = boundaries[0] if boundaries else N

    all_results = []
    residuals = []
    e_start = E_INITIAL
    max_checks = {"balance": 0.0, "storage": 0.0, "min_energy": E_MAX, "max_energy": E_MIN}

    for day in range(len(dates)):
        p_t = prices4[day]
        load0 = load_forecast_10min(loads, day, prior_load)
        pv = {0: hourly_forecast_to_10min(forecasts[day, 0], 0)}
        for h in adjust_issues:
            anchor = pvs[day][h * 6 - 1]  # 发布前最后一个实际光伏值（kWh/10min）
            pv[h] = hourly_forecast_to_10min(forecasts[day, h // 6], h, anchor)
        margin = residual_margin(residuals)

        B = solve_dayahead(load0 - pv[0] + margin, p_t, e_start)["buy"]

        charge = np.zeros(N)
        discharge = np.zeros(N)
        emergency = np.zeros(N)
        surplus = np.zeros(N)
        energy = np.zeros(N + 1)
        energy[0] = e_start
        A = B.copy()

        e = e_start
        prev = 0
        for b in boundaries:
            e = execute_range(A, loads[day], pvs[day], e, prev, b, charge, discharge, emergency, surplus)
            h = b // 6
            net = np.concatenate(
                [loads[day][:b] - pvs[day][:b], load0[b:] - pv[h][b:] + margin[b:]]
            )
            A[b:144] = solve_adjust(B, net, p_t, e, b)
            prev = b
        e = execute_range(A, loads[day], pvs[day], e, prev, N, charge, discharge, emergency, surplus)

        energy = _build_energy_trajectory(e_start, charge, discharge)

        delta = A - B
        dup = np.maximum(delta, 0.0)
        ddown = np.maximum(-delta, 0.0)
        plan_cost = float(np.dot(p_t, B))
        adjust_cost = float(np.dot(p_t[first_boundary:], 1.5 * dup[first_boundary:] - 0.5 * ddown[first_boundary:]))
        emergency_cost = float(np.dot(5.0 * p_t, emergency))

        checks = validate_day(charge, discharge, energy, emergency, surplus, A, loads[day], pvs[day])
        max_checks["balance"] = max(max_checks["balance"], checks["balance"])
        max_checks["storage"] = max(max_checks["storage"], checks["storage"])
        max_checks["min_energy"] = min(max_checks["min_energy"], checks["min_energy"])
        max_checks["max_energy"] = max(max_checks["max_energy"], checks["max_energy"])

        all_results.append(
            {
                "buy": B,
                "adjusted": A,
                "charge": charge,
                "discharge": discharge,
                "emergency": emergency,
                "surplus": surplus,
                "energy": energy,
                "plan_cost": plan_cost,
                "adjust_cost": adjust_cost,
                "emergency_cost": emergency_cost,
                "dup": dup,
                "ddown": ddown,
            }
        )

        # 承诺时刻残差：每段用锁定该段购电量的预报
        residual = np.zeros(N)
        commit_points = [0] + boundaries
        for i, b in enumerate(commit_points):
            b_next = commit_points[i + 1] if i + 1 < len(commit_points) else N
            ch = b // 6
            residual[b:b_next] = (loads[day] - pvs[day])[b:b_next] - (load0 - pv[ch])[b:b_next]
        residuals.append(residual)
        e_start = float(energy[-1])

    report = all_results[REPORT_START:]
    report_dates = dates[REPORT_START:]

    plan_kwh = sum(np.sum(r["buy"]) for r in report)
    adjusted_kwh = sum(np.sum(r["adjusted"]) for r in report)
    dup_kwh = sum(np.sum(r["dup"]) for r in report)
    ddown_kwh = sum(np.sum(r["ddown"]) for r in report)
    emergency_kwh = sum(np.sum(r["emergency"]) for r in report)
    plan_cost = sum(r["plan_cost"] for r in report)
    adjust_cost = sum(r["adjust_cost"] for r in report)
    emergency_cost = sum(r["emergency_cost"] for r in report)
    total_cost = plan_cost + adjust_cost + emergency_cost
    metrics = {
        "report_days": len(report),
        "planned_purchase": plan_kwh,
        "adjusted_purchase": adjusted_kwh,
        "up_adjust": dup_kwh,
        "down_adjust": ddown_kwh,
        "emergency_purchase": emergency_kwh,
        "planned_cost": plan_cost,
        "adjust_cost": adjust_cost,
        "emergency_cost": emergency_cost,
        "total_cost": total_cost,
        "emergency_days": sum(np.sum(r["emergency"]) > 0.01 for r in report),
        "adjust_days": sum(np.sum(r["dup"] + r["ddown"]) > 0.01 for r in report),
        "ending_energy": float(report[-1]["energy"][-1]),
        "max_balance_residual": max_checks["balance"],
        "max_storage_residual": max_checks["storage"],
        "global_min_energy": max_checks["min_energy"],
        "global_max_energy": max_checks["max_energy"],
        "simultaneous": sum(int(np.sum((r["charge"] > TOL) & (r["discharge"] > TOL))) for r in report),
    }
    write_output_4_3(report_dates, report)
    return metrics


def _fmt_metrics(tag, m):
    lines = [
        f"{tag} MODEL METRICS (波动电价)",
        f"report_days={m['report_days']}",
        f"planned_purchase={m['planned_purchase']:.4f} kWh",
        f"emergency_purchase={m['emergency_purchase']:.4f} kWh",
    ]
    if "adjusted_purchase" in m:
        lines += [
            f"adjusted_purchase={m['adjusted_purchase']:.4f} kWh",
            f"up_adjust={m['up_adjust']:.4f} kWh",
            f"down_adjust={m['down_adjust']:.4f} kWh",
        ]
    lines += [
        f"planned_cost={m['planned_cost']:.4f} yuan",
    ]
    if "adjust_cost" in m:
        lines.append(f"adjust_cost={m['adjust_cost']:.4f} yuan")
    lines += [
        f"emergency_cost={m['emergency_cost']:.4f} yuan",
        f"total_cost={m['total_cost']:.4f} yuan",
        f"emergency_days={m['emergency_days']}/{m['report_days']}",
        f"ending_energy={m['ending_energy']:.4f} kWh",
        f"max_balance_residual={m['max_balance_residual']:.8g}",
        f"max_storage_residual={m['max_storage_residual']:.8g}",
        f"global_min_energy={m['global_min_energy']:.4f} kWh",
        f"global_max_energy={m['global_max_energy']:.4f} kWh",
        f"simultaneous={m['simultaneous']}",
    ]
    return "\n".join(lines)


def main():
    prices4 = load_prices4()
    print("加载附件4波动电价：shape =", prices4.shape, "| 日均价区间 [%.3f, %.3f] 元/kWh" % (prices4.mean(axis=1).min(), prices4.mean(axis=1).max()))

    m2 = run_q2(prices4)
    print("\n" + _fmt_metrics("Q4-2", m2))
    print("output=%s" % OUTPUT4_2)
    SUMMARY4_2.write_text(_fmt_metrics("Q4-2", m2) + "\noutput=%s\n" % OUTPUT4_2, encoding="utf-8")

    m3 = run_q3(prices4)
    print("\n" + _fmt_metrics("Q4-3", m3))
    print("output=%s" % OUTPUT4_3)
    SUMMARY4_3.write_text(_fmt_metrics("Q4-3", m3) + "\noutput=%s\n" % OUTPUT4_3, encoding="utf-8")

    print("\n=== 固定电价(问题3) vs 波动电价(4-3) ===")
    print(f"固定电价 total = 13,727,335 元   vs   波动电价 total = {m3['total_cost']:.0f} 元")


if __name__ == "__main__":
    main()

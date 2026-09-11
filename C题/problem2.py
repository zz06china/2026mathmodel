# -*- coding: utf-8 -*-
"""
问题2：微网计划购电策略 —— 两阶段(计划-执行)线性规划
主模型：LP

框架：
  每天 0:00 制定"计划购电策略"，只能用当时信息(电价 + 当前SOC + 历史真实数据):
    - 用"持续性预测"(预测=昨天实际)作为当天计划依据
  当天执行/回测时用真实负载/光伏(附件2):
    - 计划只锁定"购电量 b_t"(这是0:00承诺)
    - 储能充放电在真实数据下重优化, 供电缺口才按5倍价紧急购电
  SOC 跨天连续结转, 年底(12.31 24:00)自由(不强制=6000)

每天两个 LP：
  [计划LP] 用预测(Lhat,Ghat) -> 决定 b_t (承诺购电量)
  [执行LP] 用真实(L_act,G_act), b_t 固定 -> 重优化 c,d,e (最小化紧急购电费)
"""
import openpyxl
import pulp

DT = 1 / 6.0
T = 144
ETA = 0.9
E_MIN, E_MAX = 1200.0, 10800.0
C_MAX = 5000.0 * DT
E_INIT = 6000.0

# ---------- 1. 读数据 ----------
wb1 = openpyxl.load_workbook("附件/附件1.xlsx", data_only=True)
rows1 = list(wb1["Sheet1"].iter_rows(values_only=True))[1:]
prices = [float(r[1]) for r in rows1]

wb2 = openpyxl.load_workbook("附件/附件2.xlsx", data_only=True, read_only=True)
load_sheet = wb2[wb2.sheetnames[0]]
pv_sheet = wb2[wb2.sheetnames[1]]
load_all, pv_all = [], []
for row in load_sheet.iter_rows(values_only=True):
    vals = list(row)
    if isinstance(vals[0], str):
        continue
    load_all.append([float(v) * DT for v in vals[1:]])
for row in pv_sheet.iter_rows(values_only=True):
    vals = list(row)
    if isinstance(vals[0], str):
        continue
    pv_all.append([float(v) * DT for v in vals[1:]])
wb2.close()
NDAYS = len(load_all)
print(f"载入 {NDAYS} 天数据, 每天 {T} 时段")


# ---------- 2. 计划LP: 用预测决定承诺购电量 b_t ----------
def plan_day(Lhat, Ghat, E_start):
    prob = pulp.LpProblem("plan", pulp.LpMinimize)
    b = [pulp.LpVariable(f"b{t}", 0) for t in range(T)]
    c = [pulp.LpVariable(f"c{t}", 0) for t in range(T)]
    d = [pulp.LpVariable(f"d{t}", 0) for t in range(T)]
    s = [pulp.LpVariable(f"s{t}", 0) for t in range(T)]   # 弃光/弃电
    E = [pulp.LpVariable(f"E{t}") for t in range(T)]
    prob += pulp.lpSum(prices[t] * b[t] for t in range(T))
    for t in range(T):
        prob += b[t] + Ghat[t] + d[t] == Lhat[t] + c[t] + s[t]
        Ep = E_start if t == 0 else E[t - 1]
        prob += E[t] == Ep + ETA * c[t] - d[t] / ETA
        prob += E[t] >= E_MIN
        prob += E[t] <= E_MAX
        prob += c[t] <= C_MAX
        prob += d[t] <= C_MAX
    prob.solve(pulp.PULP_CBC_CMD(msg=0))
    b_v = [pulp.value(b[t]) for t in range(T)]
    return b_v


# ---------- 3. 执行LP: 真实数据下重优化储能, 最小化紧急购电费 ----------
def execute_day(b, L_act, G_act, E_start):
    prob = pulp.LpProblem("exec", pulp.LpMinimize)
    c = [pulp.LpVariable(f"c{t}", 0) for t in range(T)]
    d = [pulp.LpVariable(f"d{t}", 0) for t in range(T)]
    e = [pulp.LpVariable(f"e{t}", 0) for t in range(T)]
    s = [pulp.LpVariable(f"s{t}", 0) for t in range(T)]   # 弃电/弃光(含超购盈余)
    E = [pulp.LpVariable(f"E{t}") for t in range(T)]
    prob += pulp.lpSum(5.0 * prices[t] * e[t] for t in range(T))
    for t in range(T):
        prob += b[t] + G_act[t] + d[t] + e[t] == L_act[t] + c[t] + s[t]
        Ep = E_start if t == 0 else E[t - 1]
        prob += E[t] == Ep + ETA * c[t] - d[t] / ETA
        prob += E[t] >= E_MIN
        prob += E[t] <= E_MAX
        prob += c[t] <= C_MAX
        prob += d[t] <= C_MAX
    prob.solve(pulp.PULP_CBC_CMD(msg=0))
    c_v = [pulp.value(c[t]) for t in range(T)]
    d_v = [pulp.value(d[t]) for t in range(T)]
    e_v = [pulp.value(e[t]) for t in range(T)]
    s_v = [pulp.value(s[t]) for t in range(T)]
    E_end = pulp.value(E[T - 1])
    return c_v, d_v, e_v, s_v, E_end


# ---------- 4. 逐日滚动 ----------
E_carry = E_INIT
total_plan_cost = 0.0
total_emerg_cost = 0.0
total_emerg_kwh = 0.0
total_spill_kwh = 0.0
daily = []

for day in range(NDAYS):
    # 0:00 计划: 持续性预测(预测=昨天实际); 第一天无昨天, 用当天实际
    Lhat = load_all[day - 1] if day >= 1 else load_all[0]
    Ghat = pv_all[day - 1] if day >= 1 else pv_all[0]
    b = plan_day(Lhat, Ghat, E_carry)

    # 当天执行/回测: 真实数据, b 固定
    L_act, G_act = load_all[day], pv_all[day]
    c, d, e, s, E_end = execute_day(b, L_act, G_act, E_carry)

    plan_cost = sum(prices[t] * b[t] for t in range(T))
    emerg_cost = sum(5.0 * prices[t] * e[t] for t in range(T))
    emerg_kwh = sum(e)
    spill_kwh = sum(s)
    total_plan_cost += plan_cost
    total_emerg_cost += emerg_cost
    total_emerg_kwh += emerg_kwh
    total_spill_kwh += spill_kwh
    daily.append({"b": b, "c": c, "d": d, "e": e, "s": s, "E_start": E_carry, "E_end": E_end})
    E_carry = E_end

print("\n===== 全年结果 =====")
print(f"计划购电费(元): {total_plan_cost:,.2f}")
print(f"紧急购电费(元): {total_emerg_cost:,.2f}")
print(f"全年总购电费(元): {total_plan_cost + total_emerg_cost:,.2f}")
print(f"全年计划购电量(kWh): {sum(sum(d['b']) for d in daily):,.2f}")
print(f"全年紧急购电量(kWh): {total_emerg_kwh:,.2f}")
print(f"全年弃电量(kWh): {total_spill_kwh:,.2f}")
print(f"年末 SOC (12.31 24:00): {E_carry:.2f} kWh")

# 紧急购电发生天数统计
emerg_days = sum(1 for d in daily if sum(d["e"]) > 0.01)
print(f"发生紧急购电的天数: {emerg_days}/{NDAYS}")


# ---------- 5. 写 result2.xlsx (2025.2.1 ~ 2025.12.31, 共334天) ----------
from datetime import datetime, time, timedelta

REPORT_START = 31   # 附件2 第32天 = 2025-02-01
REPORT_END = 364    # 第365天 = 2025-12-31
ND = REPORT_END - REPORT_START + 1   # 334

wb_out = openpyxl.load_workbook("附件/附件5/result2.xlsx")

# (1) 计划购电量: 每天144时段购电量 + 全天购电量 + 全天购电费
ws1 = wb_out["计划购电量"]
for i in range(ND):
    d = daily[REPORT_START + i]
    b = d["b"]
    row = i + 2
    for t in range(T):
        ws1.cell(row=row, column=t + 2, value=round(b[t], 4))
    ws1.cell(row=row, column=146, value=round(sum(b), 4))
    ws1.cell(row=row, column=147, value=round(sum(prices[t] * b[t] for t in range(T)), 4))

# (2) 充放电量: 每天6个4小时时段充/放电量 + 0:00/24:00储电量
ws2 = wb_out["充放电量"]
ws2.delete_rows(2, ws2.max_row)   # 清掉模板样例数据
block_labels = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
                "12:00-16:00", "16:00-20:00", "20:00-24:00"]
r = 2
for i in range(ND):
    d = daily[REPORT_START + i]
    date = datetime(2025, 1, 1) + timedelta(days=REPORT_START + i)
    for k in range(6):
        i0, i1 = k * 24, (k + 1) * 24
        if k == 0:
            ws2.cell(row=r, column=1, value=date)
            ws2.cell(row=r, column=5, value=time(0, 0))
            ws2.cell(row=r, column=6, value=round(d["E_start"], 4))
        if k == 1:
            ws2.cell(row=r, column=5, value="24:00")
            ws2.cell(row=r, column=6, value=round(d["E_end"], 4))
        ws2.cell(row=r, column=2, value=block_labels[k])
        ws2.cell(row=r, column=3, value=round(sum(d["c"][i0:i1]), 4))
        ws2.cell(row=r, column=4, value=round(sum(d["d"][i0:i1]), 4))
        r += 1

# (3) 紧急购电量: 每天把连续紧急购电时段合并成时间段
ws3 = wb_out["紧急购电量"]
ws3.delete_rows(2, ws3.max_row)


def slot_label(t0, t1):
    def hhmm(m):
        m %= 1440
        return f"{m // 60:02d}:{m % 60:02d}"
    start = 10 + t0 * 10
    end = 10 + (t1 + 1) * 10
    e = hhmm(end) + ("+1" if end > 1440 else "")
    return f"{hhmm(start)}-{e}"


r = 2
for i in range(ND):
    d = daily[REPORT_START + i]
    e = d["e"]
    date = datetime(2025, 1, 1) + timedelta(days=REPORT_START + i)
    wrote_date = False
    t = 0
    while t < T:
        if e[t] > 0.01:
            t0 = t
            while t < T and e[t] > 0.01:
                t += 1
            t1 = t - 1
            ws3.cell(row=r, column=1, value=date if not wrote_date else None)
            ws3.cell(row=r, column=2, value=slot_label(t0, t1))
            ws3.cell(row=r, column=3, value=round(sum(e[t0:t1 + 1]), 4))
            wrote_date = True
            r += 1
        else:
            t += 1

wb_out.save("result2.xlsx")
print(f"\n已保存 result2.xlsx: 计划购电量 {ND} 天, 充放电量 {ND} 天, 紧急购电量 {r - 2} 条")

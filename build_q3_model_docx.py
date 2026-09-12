import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
Q1_DIR = ROOT / "C题_Q1_最终模型"
sys.path.insert(0, str(Q1_DIR))

import build_q1_paper as paper  # noqa: E402


SOURCE = HERE / "第三问最终模型.md"
OUTPUT = HERE / "第三问最终模型_公式修正版.docx"

TABLE_CAPTIONS = {
    "4. 符号说明": "表1  主要符号说明",
    "8. 全年结果与模型比较": "表2  不同策略全年结果比较",
    "8.1 各更新时刻的边际价值": "表3  各预报更新时刻的边际价值",
    "9.1 0:00 计划购电量": "表4  指定日期的0:00计划购电结果",
    "9.2 最终调整购电量": "表5  指定日期的最终调整购电结果",
    "9.3 储能运行结果": "表6  指定日期的储能运行结果",
    "9.4 紧急购电结果": "表7  指定日期的紧急购电结果",
}


def add_inline_content(paragraph, text, default_size=12):
    """Add Markdown bold and inline LaTeX as Word-native runs/OMML."""
    token_pattern = re.compile(r"(\$[^$]+\$|\*\*.*?\*\*)")
    for token in token_pattern.split(text):
        if not token:
            continue
        if token.startswith("$") and token.endswith("$"):
            paragraph._p.append(paper.latex_omml(token[1:-1]))
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            paper.set_run_font(run, size=default_size, bold=True)
        else:
            run = paragraph.add_run(token)
            paper.set_run_font(run, size=default_size)


def add_mixed_paragraph(doc, text, first_line=True):
    paragraph = doc.add_paragraph()
    add_inline_content(paragraph, text)
    paper.style_paragraph(paragraph, first_line=first_line)
    return paragraph


def add_numbered_item(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(-24)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(f"{number}. ")
    paper.set_run_font(run)
    add_inline_content(paragraph, text)
    return paragraph


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_value(cell, value, bold=False, font_size=9.5, left=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    exact_math = re.fullmatch(r"\$([^$]+)\$", value)
    if exact_math:
        paragraph._p.append(paper.latex_omml(exact_math.group(1)))
    else:
        add_inline_content(paragraph, value, default_size=font_size)
        if bold:
            for run in paragraph.runs:
                run.bold = True


def table_widths(column_count):
    if column_count == 3:
        return [3.0, 8.5, 3.0]
    if column_count == 5:
        return [2.4, 2.7, 2.8, 2.8, 3.8]
    if column_count == 6:
        return [3.7, 2.15, 2.15, 2.15, 2.35, 2.0]
    if column_count == 9:
        return [2.0, 1.35, 1.35, 1.35, 1.35, 1.35, 1.35, 2.15, 1.95]
    return [14.5 / column_count] * column_count


def add_markdown_table(doc, headers, rows, caption, symbol_table=False):
    paper.add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    widths = table_widths(len(headers))
    compact_size = 7.3 if len(headers) >= 8 else (8.2 if len(headers) >= 5 else 9.5)

    for index, header in enumerate(headers):
        set_cell_value(table.rows[0].cells[index], header, bold=True, font_size=compact_size)
    set_repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            left = (symbol_table and index == 1) or (len(headers) <= 3 and index == 1)
            set_cell_value(cells[index], value, font_size=compact_size, left=left)

    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Cm(width)
            tc_w = row.cells[index]._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(Cm(width).twips))
            tc_w.set(qn("w:type"), "dxa")

    paper.format_three_line_table(table)
    return table


def add_title(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(text)
    paper.set_run_font(run, east_asia="黑体", size=18, bold=True)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    paper.configure_document(document)
    document.core_properties.title = "问题三最终模型（公式修正版）"
    document.core_properties.subject = "非对称结算约束下的分位数鲁棒模型预测控制"
    document.core_properties.author = "数学建模论文组"

    equation_number = 0
    current_heading = ""
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue

        if line.startswith("# "):
            add_title(document, line[2:].strip())
            current_heading = line[2:].strip()
            index += 1
            continue
        if line.startswith("## "):
            current_heading = line[3:].strip()
            paper.add_heading(document, current_heading, 1)
            index += 1
            continue
        if line.startswith("### "):
            current_heading = line[4:].strip()
            paper.add_heading(document, current_heading, 2)
            index += 1
            continue

        if line.startswith("$$"):
            formula_parts = [line[2:]]
            while not formula_parts[-1].endswith("$$"):
                index += 1
                formula_parts.append(lines[index].strip())
            formula_parts[-1] = formula_parts[-1][:-2]
            latex = " ".join(formula_parts).strip()
            equation_number += 1
            paper.add_equation(document, latex, equation_number)
            index += 1
            continue

        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            headers = split_table_row(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            caption = TABLE_CAPTIONS.get(current_heading, f"表  {current_heading}")
            add_markdown_table(
                document,
                headers,
                rows,
                caption,
                symbol_table=current_heading == "4. 符号说明",
            )
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            add_numbered_item(document, numbered.group(1), numbered.group(2))
        else:
            add_mixed_paragraph(document, line)
        index += 1

    document.save(OUTPUT)
    return OUTPUT, equation_number


if __name__ == "__main__":
    output, equations = build_document()
    print(f"saved={output}")
    print(f"equations={equations}")

